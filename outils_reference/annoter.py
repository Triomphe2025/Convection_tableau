import re, pickle, openpyxl
from collections import Counter
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from comparateur import *
from pieds import champs, pages_ref_txt, pages_xls, comparer_pieds

ROUGE = PatternFill('solid', fgColor='FF9999'); ORANGE = PatternFill('solid', fgColor='FFC000')
JAUNE = PatternFill('solid', fgColor='FFF2A8'); GRIS = PatternFill('solid', fgColor='D9D9D9')
PRIO = {'Confusion de caractère': (1, 'À corriger', ORANGE), 'Contenu différent': (1, 'À corriger (ou mise à jour volontaire ?)', ROUGE),
        "Ajouté dans l'Excel": (1, 'À corriger', ROUGE), "Manquant dans l'Excel": (1, 'À corriger', ROUGE),
        'Pied de page': (1, 'À corriger', ROUGE), "Ligne absente de l'Excel": (1, 'À ajouter', ROUGE),
        "Ligne en trop dans l'Excel": (2, 'À vérifier', ROUGE),
        'Espacement interne': (3, 'Position', JAUNE), 'Position des sous-champs': (3, 'Position', JAUNE),
        'Non vérifiable': (4, 'Info', GRIS)}

def lignes_brutes(path):
    return [l for l in open(path, encoding='utf-8', errors='replace') if not l.startswith('|')]

def generer(nom, src_xlsx, R, X, feet, ecarts, pos, orph, pieds_ec, cols, brut=None, note=''):
    wb = openpyxl.load_workbook(src_xlsx)
    ws = wb.worksheets[-1]; sheet = ws.title
    items = []
    def add(typ, page, row, col, orig, xls, info=''):
        p, action, fill = PRIO[typ]
        items.append((p, int(row) if row else 10**9, typ, action, page, row, col, orig, xls, info, fill))
    for e in ecarts: add(e['type'], e['page'], e['row'], e['col'], e['orig'], e['xls'], e['freq'])
    for e in pos: add('Position des sous-champs', e['page'], e['row'], e['col'], e['orig'], e['xls'],
                      f"colonnes de début : original {e['po']}  /  Excel {e['px']}")
    brut_norm = [set(l.split()) for l in (brut or [])]
    for o in orph:
        if o['row'] and brut:
            toks = set(o['xls'].replace('|', ' ').split())
            if toks and max((len(toks & b) / len(toks) for b in brut_norm), default=0) >= 0.6:
                add('Non vérifiable', o['page'], o['row'], 0, '', o['xls'], "ligne présente dans l'original, sur une page au format OCR (sans cadre) : non comparée")
                continue
        add(o['type'], o['page'], o['row'], 0, o['orig'], o['xls'])
    for e in pieds_ec:
        add('Pied de page', e['page'], e['rows'][-1] if e['champ'] in ('N° PLAN', 'INDICE', 'PAGE') else e['rows'][0], 1,
            f"{e['champ']} : {e['orig']}", f"{e['champ']} : {e['xls']}")
    items.sort()
    # --- annotation des cellules
    for p, _, typ, action, page, row, col, orig, xls, info, fill in items:
        if not row: continue
        c = ws.cell(row=int(row), column=col + 1)
        if type(c).__name__ == 'MergedCell':
            for rg in ws.merged_cells.ranges:
                if c.coordinate in rg: c = ws.cell(row=rg.min_row, column=rg.min_col); break
        c.fill = fill
        txt = f"{typ}\nOriginal : «{orig}»" + (f"\n{info}" if info else '')
        if c.comment: txt = c.comment.text + '\n---\n' + txt
        c.comment = Comment(txt[:1500], 'Contrôle'); c.comment.width = 380; c.comment.height = 140
    # --- feuille des ecarts
    es = wb.create_sheet('ÉCARTS', 0)
    H = ['Priorité', 'Action', 'Type', 'Page', 'Cellule', 'Colonne', 'Valeur dans ton Excel', "Valeur de l'original", 'Aide à la décision']
    es.append(H)
    for c, w in zip(es[1], [8, 22, 24, 7, 11, 13, 36, 36, 60]):
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
        es.column_dimensions[c.column_letter].width = w
    for p, _, typ, action, page, row, col, orig, xls, info, fill in items:
        addr = f"{get_column_letter(col + 1)}{row}" if row else '—'
        colname = cols[col] if typ not in ('Pied de page', "Ligne absente de l'Excel", "Ligne en trop dans l'Excel", 'Non vérifiable') else ('Pied' if typ == 'Pied de page' else 'Ligne')
        es.append([p, action, typ, page, addr, colname, xls, orig, info])
        if row:
            cell = es.cell(row=es.max_row, column=5)
            from openpyxl.worksheet.hyperlink import Hyperlink
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{addr}", display=addr); cell.font = Font(color='0563C1', underline='single')
        es.cell(row=es.max_row, column=3).fill = fill
    es.freeze_panes = 'A2'; es.auto_filter.ref = es.dimensions
    # --- resume
    rs = wb.create_sheet('RÉSUMÉ', 0)
    cnt = Counter((it[3], it[2]) for it in items)
    rs.append([f"Contrôle de {nom} — Excel corrigé comparé à l'original"]); rs['A1'].font = Font(bold=True, size=14)
    rs.append([note]); rs.append([])
    rs.append(['Action', 'Type', 'Nombre']); [setattr(c, 'font', Font(bold=True)) for c in rs[4]]
    for (a, t), n in sorted(cnt.items(), key=lambda kv: PRIO[kv[0][1]][0]): rs.append([a, t, n])
    rs.append([])
    for l in ["Mode d'emploi : filtre la feuille ÉCARTS sur Priorité = 1, clique sur l'adresse de cellule pour y aller.",
              "Dans la feuille Listing, les cellules en cause sont colorées et portent un commentaire avec la valeur de l'original.",
              "Orange = caractère confondu (erreur de lecture probable). Rouge = contenu, ligne ou pied différent. Jaune = position des sous-champs. Gris = non vérifiable.",
              "Colonne « Aide à la décision » : nombre d'occurrences de chaque forme dans tout l'original et dans tout ton Excel.",
              "Copie de travail : le logo .wmf de la page de garde n'est pas repris par la bibliothèque utilisée. Corrige de préférence dans ton fichier d'origine."]:
        rs.append([l])
    rs.column_dimensions['A'].width = 40; rs.column_dimensions['B'].width = 32
    out = f'{nom}_ANNOTE.xlsx'; wb.save(out)
    return out, cnt
