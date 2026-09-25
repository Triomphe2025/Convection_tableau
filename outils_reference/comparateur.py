"""Comparaison original <-> Excel corrige, avec adresse de cellule.
Pas d'IA : alignement de sequences + regles."""
import re, difflib, json, unicodedata
from collections import Counter
import openpyxl
from openpyxl.utils import get_column_letter

COLS4 = 4
FOOT_RE = re.compile(r"N°\s?PLAN|NO\s?PLAN|P\.?E\.?T\.?\s*:|CABLE\s*:|TYPE\s*:|INDICE\s*:|PAGE\s*:|M A T R A|^MATRA|SIEMENS", re.I)

def norm(s):  # contenu : espaces multiples reduits
    return ' '.join((s or '').split())

def offsets(s):
    s = (s or '').rstrip()
    lead = len(s) - len(s.lstrip())
    return tuple(m.start() - lead for m in re.finditer(r'\S+', s))

# ---------------------------------------------------------------- Excel
def lire_excel(path, entete_mots):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[-1]
    rows, pending, feet = [], [], []
    for r in ws.iter_rows(min_row=1, max_col=COLS4):
        v = ['' if c.value is None else str(c.value) for c in r]
        j = ' '.join(v)
        if FOOT_RE.search(j):
            m = re.search(r'PAGE\s*:?\s*(\w+)', j, re.I)
            if m:
                for x in pending: x['page'] = m.group(1)
                pending = []
            feet.append((r[0].row, j))
            continue
        if norm(v[0]).upper() in entete_mots or not any(x.strip() for x in v):
            continue
        d = {'row': r[0].row, 'cells': v, 'page': '?'}
        rows.append(d); pending.append(d)
    return wb, ws, rows, feet

# ---------------------------------------------------------------- references
def lire_txt(path, entete_mots):
    rows, page_rows, pageno = [], [], 0
    for l in open(path, encoding='utf-8', errors='replace'):
        l = l.rstrip('\r\n')
        m = re.search(r'PAGE\s*:\s*(\w+)', l)
        if m:
            for x in page_rows: x['page'] = m.group(1)
            page_rows = []
        if not l.startswith('|') or l.count('|') < 5: continue
        c = [x.strip() for x in l.split('|')[1:-1]]
        if len(c) != 4 or c[0].upper() in entete_mots: continue
        if not any(c) or set(''.join(c)) <= set('-_'): continue
        if FOOT_RE.search(' '.join(c)): continue
        d = {'cells': c, 'page': '?', 'exact': False}
        rows.append(d); page_rows.append(d)
    return rows

def lire_pdf137(path):
    import pymupdf
    from grille import grille
    doc = pymupdf.open(path); rows = []
    for i, p in enumerate(doc):
        g = grille(p)
        if not g: continue
        pg = None
        for l in g:
            m = re.search(r'PAGE\s*:?\s*(\w+)\s*\|?\s*$', l)
            if m: pg = m.group(1)
        for l in g:
            if FOOT_RE.search(l): continue
            if l.count('|') >= 5:
                segs = l.split('|')[1:-1]
                if len(segs) != 4: continue
                c = [s[1:].rstrip() if s.startswith(' ') else s.rstrip() for s in segs]
            elif re.match(r'^\S.{20,}\d{4}[A-Z]', l):          # format TP2 sans barres
                c = [l[:22].rstrip(), l[22:32].strip(), l[32:51].strip(), l[51:].strip()]
                c[2] = l[32:51].strip(); c[2] = l[l.find(c[2], 32):51].rstrip() if c[2] else ''
            else:
                continue
            if not any(x.strip() for x in c) or norm(c[0]).upper().startswith(('TENANT', '°')):
                continue
            rows.append({'cells': c, 'page': pg or str(i + 1), 'exact': True})
    return rows

# ---------------------------------------------------------------- alignement
def _sig(r): return ' | '.join(norm(x) for x in r['cells'])

def _nw(A, B):
    """alignement optimal (Needleman-Wunsch) de deux petits blocs de lignes"""
    n, m = len(A), len(B)
    if n * m > 250000:
        return [(i, i if i < m else None) for i in range(n)] + [(None, j) for j in range(n, m)]
    S = [[difflib.SequenceMatcher(None, a, b, autojunk=False).ratio() for b in B] for a in A]
    gap = -0.35
    F = [[0.0] * (m + 1) for _ in range(n + 1)]
    for i in range(1, n + 1): F[i][0] = i * gap
    for j in range(1, m + 1): F[0][j] = j * gap
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            F[i][j] = max(F[i-1][j-1] + S[i-1][j-1] - 0.5, F[i-1][j] + gap, F[i][j-1] + gap)
    out, i, j = [], n, m
    while i or j:
        if i and j and abs(F[i][j] - (F[i-1][j-1] + S[i-1][j-1] - 0.5)) < 1e-9:
            out.append((i-1, j-1)); i -= 1; j -= 1
        elif i and abs(F[i][j] - (F[i-1][j] + gap)) < 1e-9:
            out.append((i-1, None)); i -= 1
        else:
            out.append((None, j-1)); j -= 1
    return out[::-1]

def aligner(R, X):
    a = [_sig(r) for r in R]; b = [_sig(r) for r in X]
    pairs = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op == 'equal':
            pairs += list(zip(range(i1, i2), range(j1, j2)))
        else:
            for i, j in _nw(a[i1:i2], b[j1:j2]):
                pairs.append((None if i is None else i1 + i, None if j is None else j1 + j))
    return pairs

# ---------------------------------------------------------------- classement
CONF = set(map(frozenset, [('1','I'),('1','L'),('0','O'),('0','D'),('5','S'),('8','B'),('2','Z'),('6','G'),
       ('P','F'),('D','O'),('I','L'),(' ','_'),("'",'"'),('R','_'),('M','N'),('E','F'),('C','G'),('H','K'),
       ('U','V'),('V','Y'),('1','7'),('R','K'),('T','I')]))

def classe(a, b):
    if a.replace(' ', '') == b.replace(' ', ''):
        return 'Espacement interne'
    ch = [(a[i1:i2], b[j1:j2]) for t, i1, i2, j1, j2 in
          difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes() if t != 'equal']
    if len(ch) <= 2 and all(len(x) <= 1 and len(y) <= 1 and (x == '' or y == '' or frozenset((x, y)) in CONF) for x, y in ch):
        return 'Confusion de caractère'
    if not a: return 'Ajouté dans l\'Excel'
    if not b: return 'Manquant dans l\'Excel'
    return 'Contenu différent'

def comparer(R, X, noms_cols):
    toks = lambda rows: Counter(t for r in rows for c in r['cells'] for t in norm(c).split())
    fr, fx = toks(R), toks(X)
    ecarts, pos, orph = [], [], []
    for i, j in aligner(R, X):
        if j is None:
            orph.append({'type': 'Ligne absente de l\'Excel', 'page': R[i]['page'],
                         'row': None, 'orig': ' | '.join(norm(x) for x in R[i]['cells']), 'xls': ''})
            continue
        if i is None:
            orph.append({'type': 'Ligne en trop dans l\'Excel', 'page': X[j]['page'], 'row': X[j]['row'],
                         'orig': '', 'xls': ' | '.join(norm(x) for x in X[j]['cells'])})
            continue
        r, x = R[i], X[j]
        for k in range(4):
            a, b = norm(r['cells'][k]), norm(x['cells'][k])
            if a != b:
                cl = classe(a, b)
                dif = [t for t in b.split() if t not in a.split()]
                ref = [t for t in a.split() if t not in b.split()]
                ecarts.append({'type': cl, 'page': x['page'], 'row': x['row'], 'col': k, 'colname': noms_cols[k],
                               'orig': r['cells'][k].rstrip() if r['exact'] else a, 'xls': x['cells'][k].rstrip(),
                               'freq': '; '.join(f"{t}: orig {fr[t]} / xls {fx[t]}" for t in (ref + dif)[:4])})
            elif r['exact'] and offsets(r['cells'][k]) != offsets(x['cells'][k]):
                pos.append({'page': x['page'], 'row': x['row'], 'col': k, 'colname': noms_cols[k],
                            'orig': r['cells'][k].rstrip(), 'xls': x['cells'][k].rstrip(),
                            'po': offsets(r['cells'][k]), 'px': offsets(x['cells'][k])})
    return ecarts, pos, orph
