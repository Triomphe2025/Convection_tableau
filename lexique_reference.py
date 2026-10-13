"""
Lexique de référence de la règle de cohérence (étape 9) : codes tirés des Excel déjà
corrigés par l'utilisateur (QUALITE_LEXIQUE_DOSSIER, feuille QUALITE_LEXIQUE_FEUILLE).

Ces Excel contiennent eux-mêmes des erreurs : le lexique ne sert qu'à signaler et proposer,
jamais à modifier. Pour mesurer sans tricher, un document peut être exclu (`exclure`).
Un Excel corrigé s'ajoute au lexique par `ajouter_au_lexique` (bouton « Ajouter au lexique »).
Au build, les classeurs de référence sont compilés (`compiler_lexique`, formes seules) dans un
fichier embarqué dans l'exe ; à l'exécution, `lexique_complet` réunit ces formes et celles des
classeurs ajoutés par l'utilisateur.
"""
import hashlib
import json
import shutil
from datetime import date
from pathlib import Path
from typing import Dict, Optional, Set, Tuple

from config import Config
from controle_qualite import formes


def _entetes_connus() -> Set[str]:
    from template import TemplateManager
    gestion = TemplateManager()
    return {c.upper() for nom in gestion.names() for c in gestion.get(nom).columns}


def formes_classeur(chemin: Path, feuille: Optional[str] = None,
                    entetes: Optional[Set[str]] = None) -> Set[str]:
    """Formes des cellules de la feuille d'un classeur (en-têtes et pieds écartés) ; vide
    si la feuille manque."""
    import openpyxl
    feuille = feuille or Config.QUALITE_LEXIQUE_FEUILLE
    entetes = _entetes_connus() if entetes is None else entetes
    trouvees: Set[str] = set()
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    try:
        if feuille not in wb.sheetnames:
            return trouvees
        for ligne in wb[feuille].iter_rows(values_only=True):
            valeurs = [str(v) for v in ligne if v is not None and str(v).strip()]
            if valeurs and all(v.strip().upper() in entetes for v in valeurs):
                continue
            for v in valeurs:
                # Pied de page (« P.E.T. : … », « NO PLAN : … ») : pas un code du tableau.
                if ':' not in v:
                    trouvees.update(formes(v))
    finally:
        wb.close()
    return trouvees


def charger_lexique(dossier: Path, exclure: Optional[str] = None,
                    feuille: Optional[str] = None) -> Set[str]:
    """Formes des cellules des feuilles « Listing » (en-têtes et pieds de page écartés)."""
    dossier = Path(dossier)
    if not dossier.is_dir():
        return set()
    entetes = _entetes_connus()
    lexique: Set[str] = set()
    for chemin in sorted(dossier.glob('*.xlsx')):
        if exclure and exclure.upper() in chemin.stem.upper().replace(' ', ''):
            continue
        lexique |= formes_classeur(chemin, feuille, entetes)
    return lexique


def ajouter_au_lexique(source: Path, dossier: Optional[Path] = None,
                       remplacer: bool = False) -> Dict:
    """Copie un Excel corrigé dans le dossier du lexique ; formes qu'il ajoute au lexique.

    Refusé si ce n'est pas un .xlsx lisible avec une feuille QUALITE_LEXIQUE_FEUILLE
    (ValueError), ou si un fichier du même nom y est déjà et remplacer est faux
    (FileExistsError). Le dictionnaire de données n'est jamais touché.
    """
    import openpyxl
    source = Path(source)
    dossier = Path(dossier or Config.QUALITE_LEXIQUE_DOSSIER)
    feuille = Config.QUALITE_LEXIQUE_FEUILLE
    if source.suffix.lower() != '.xlsx':
        raise ValueError(f"« {source.name} » n'est pas un classeur .xlsx")
    try:
        wb = openpyxl.load_workbook(source, read_only=True)
        feuilles = list(wb.sheetnames)
        wb.close()
    except Exception as exc:
        raise ValueError(f"« {source.name} » illisible : {exc}") from exc
    if feuille not in feuilles:
        raise ValueError(f"« {source.name} » n'a pas de feuille « {feuille} » "
                         f"(feuilles : {', '.join(feuilles)})")
    cible = dossier / source.name
    if cible.exists() and not remplacer:
        raise FileExistsError(str(cible))
    # Lexique sans l'ancienne version du même fichier : on compte ce que celui-ci apporte.
    connues = {f for chemin in dossier.glob('*.xlsx') if chemin.name != source.name
               for f in formes_classeur(chemin, feuille)} if dossier.is_dir() else set()
    apportees = formes_classeur(source, feuille)
    dossier.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, cible)
    return {'chemin': cible, 'formes': len(apportees), 'ajoutees': len(apportees - connues),
            'remplace': remplacer}


def compiler_lexique(dossier: Path, sortie: Path, feuille: Optional[str] = None) -> Dict:
    """Formes des classeurs de référence, par classeur (nom, empreinte), dans un fichier JSON
    embarqué dans l'exe : l'exe n'a pas besoin des Excel pour avoir le lexique."""
    dossier, sortie = Path(dossier), Path(sortie)
    entetes = _entetes_connus()
    sources = {}
    for chemin in sorted(dossier.glob('*.xlsx')):
        sources[chemin.name] = {
            'empreinte': hashlib.sha256(chemin.read_bytes()).hexdigest(),
            'formes': sorted(formes_classeur(chemin, feuille, entetes)),
        }
    contenu = {'compile_le': date.today().isoformat(), 'sources': sources}
    sortie.parent.mkdir(parents=True, exist_ok=True)
    sortie.write_text(json.dumps(contenu, ensure_ascii=False, indent=1), encoding='utf-8')
    return contenu


def lire_lexique_compile(chemin: Path) -> Dict:
    """Contenu d'un lexique compilé ; {} si le fichier manque ou est illisible."""
    try:
        return json.loads(Path(chemin).read_text('utf-8'))
    except (OSError, ValueError):
        return {}


def lexique_complet(compile: Path, dossier_ajouts: Path, reference: Optional[Path] = None,
                    feuille: Optional[str] = None) -> Tuple[Set[str], str]:
    """(formes, bilan lisible) : lexique compilé (sinon, en développement, les classeurs de
    référence lus directement) + classeurs ajoutés par l'utilisateur."""
    contenu = lire_lexique_compile(compile)
    if contenu.get('sources'):
        integrees = {f for s in contenu['sources'].values() for f in s['formes']}
        origine = (f"{len(integrees)} formes intégrées ({len(contenu['sources'])} classeurs, "
                   f"compilé le {contenu.get('compile_le', '?')})")
    elif reference is not None and Path(reference).is_dir():
        integrees = charger_lexique(reference, feuille=feuille)
        origine = f"{len(integrees)} formes des classeurs de référence (non compilés)"
    else:
        integrees, origine = set(), "aucune forme intégrée"
    ajoutees = charger_lexique(dossier_ajouts, feuille=feuille)
    return integrees | ajoutees, (f"{origine} + {len(ajoutees - integrees)} ajoutée(s) "
                                  f"(dossier {dossier_ajouts})")


if __name__ == '__main__':
    # Build (build_exe.bat) : python lexique_reference.py --compiler
    import sys
    if sys.argv[1:] == ['--compiler']:
        from config import chemin_ressource
        bilan = compiler_lexique(Config.QUALITE_LEXIQUE_REFERENCE,
                                 chemin_ressource(Config.QUALITE_LEXIQUE_COMPILE))
        formes_total = len({f for s in bilan['sources'].values() for f in s['formes']})
        print(f"Lexique compilé : {formes_total} formes, {len(bilan['sources'])} classeurs "
              f"({', '.join(bilan['sources'])})")
        sys.exit(0 if bilan['sources'] else 1)
    print("Usage : python lexique_reference.py --compiler")
    sys.exit(2)
