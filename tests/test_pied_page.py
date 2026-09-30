"""
Pied de page : tout ce qui est lu est gardé, rien n'est inventé ni remplacé.

pied_page.py (paires LIBELLÉ : valeur, COMPLEMENT, révisions), lecture en grille,
bloc PIED_BRUT recopié par Claude, rendu brut dans l'Excel, contrôles en alerte.

Lancement :
    env\\Scripts\\python.exe -m pytest tests/test_pied_page.py -v
"""
import contextlib
import io
import json
import re
import tempfile
import unittest
from pathlib import Path

import fitz
import openpyxl

import pdf_extractor as pe
from claude_ocr import LogReplayer, _build_prompt, _parse_pipe_response
from generer_classeur import alertes_pied, generer_excel
from mesure_precision import mesurer
from mesurer_precision import lire_verite_excel, lire_xlsx
from pied_page import (analyser_pied, cle_libelle, indices_revisions, mots_decor,
                       nettoyer_lignes_pied)
from template import TemplateManager

FIXTURES = Path(__file__).parent / 'fixtures'
EXTRAIT = FIXTURES / '223111PE011_extrait_10pages.pdf'
VERITE = FIXTURES / '223111PE011_extrait_verite.xlsx'
PE137 = FIXTURES / '223400PE137.pdf'


def _tpl(nom='REPARTITEUR 2'):
    return TemplateManager().get(nom)


def _page(meta, page='7', image='page_7.png', tpl=None):
    tpl = tpl or _tpl()
    return {
        'success': True, 'headers': list(tpl.columns), 'metadata': meta,
        'image_path': image, 'detection_method': 'claude-vision',
        'rows': [{'type': 'data', 'cells': ['G', 'PH Q 01', 'TEL', 'PH A 01'],
                  'confidence': [100] * 4}],
    }


def _excel(resultats, tpl):
    """(messages du journal, texte de chaque cellule) du classeur produit."""
    messages = []
    with tempfile.TemporaryDirectory() as tmp:
        sortie = Path(tmp) / 's.xlsx'
        with contextlib.redirect_stdout(io.StringIO()):
            generer_excel(resultats, pe.PdfTableExtractor(tpl), sortie,
                          on_log=messages.append)
        cellules = [str(v) for ligne in openpyxl.load_workbook(sortie).worksheets[0]
                    .iter_rows(values_only=True) for v in ligne if v is not None]
    return messages, cellules


# ── pied_page.py ──────────────────────────────────────────────────────

class TestAnalyserPied(unittest.TestCase):

    def test_valeur_un_mot_et_complement(self):
        self.assertEqual(analyser_pied(['TYPE : 30P887 CORDON TYPE 40']),
                         {'TYPE': '30P887', 'COMPLEMENT': 'CORDON TYPE 40'})

    def test_valeur_avec_point_barre_et_tiret(self):
        self.assertEqual(analyser_pied(['TYPE : 7P.279      6/10'])['TYPE'], '7P.279')
        self.assertEqual(analyser_pied(['TYPE : 3/RO2V 2.5mm2']),
                         {'TYPE': '3/RO2V', 'COMPLEMENT': '2.5mm2'})

    def test_libelle_multi_mots_garde_son_segment(self):
        meta = analyser_pied(['P.E.T. : EPEULE MONTESQUIEU   BORNIER : B702A'])
        self.assertEqual(meta, {'PET': 'EPEULE MONTESQUIEU', 'BORNIER': 'B702A'})

    def test_diametre_hors_page_est_un_complement(self):
        meta = analyser_pied(['N° PLAN : 223111PE011 12/10'])
        self.assertEqual(meta, {'NO_PLAN': '223111PE011', 'COMPLEMENT': '12/10'})

    def test_compteur_apres_page_ou_folio(self):
        self.assertEqual(analyser_pied(['FOLIO : 3/10']), {'FOLIO': '3/10'})
        self.assertEqual(analyser_pied(['PAGE 32'])['PAGE'], '32')

    def test_texte_libre_sans_libelle(self):
        self.assertEqual(analyser_pied(['REF CE 8707905']), {'COMPLEMENT': 'REF CE 8707905'})

    def test_libelle_inconnu_conserve(self):
        self.assertEqual(analyser_pied(['ARMOIRE : AR 12  PAGE : 3'])['ARMOIRE'], 'AR 12')

    def test_fin_de_valeur_pas_prise_pour_libelle(self):
        self.assertEqual(analyser_pied(['CABLE : WPHR/TELMG TYPE : 7P.279'])['CABLE'],
                         'WPHR/TELMG')

    def test_libelle_repete_premiere_lecture_gardee(self):
        self.assertEqual(analyser_pied(['PAGE : 3', 'PAGE : 4'])['PAGE'], '3')

    def test_decor_du_modele_ni_valeur_ni_complement(self):
        meta = analyser_pied(['P.E.T. : GRAND-BUT      JARRETIERAGE'], decor={'JARRETIERAGE'})
        self.assertEqual(meta, {'PET': 'GRAND-BUT'})

    def test_libelle_vide_absent(self):
        self.assertEqual(analyser_pied(['TYPE       :']), {})


class TestNettoyerLignesPied(unittest.TestCase):

    def test_cadre_logo_et_separateurs_retires(self):
        lignes = ['|               |    P.E.T. : GRAND-BUT        JARRETIERAGE   |',
                  '|   M A T R A   |-------------------------------------------|',
                  '|               |  N° PLAN  : 223400PE137  |  INDICE : 10  |   PAGE : 1  |',
                  '|_______________________________________________________________|']
        self.assertEqual(nettoyer_lignes_pied(lignes, 'M  T  I'), [
            'P.E.T. : GRAND-BUT        JARRETIERAGE',
            'N° PLAN  : 223400PE137  |  INDICE : 10  |   PAGE : 1'])

    def test_marque_seule_retiree_texte_chiffre_garde(self):
        self.assertEqual(nettoyer_lignes_pied(['SIEMENS', 'MATRA', 'REF CE 8707905'], ''),
                         ['REF CE 8707905'])

    def test_espaces_interieurs_conserves(self):
        self.assertEqual(nettoyer_lignes_pied(['   TYPE : 7P.279      6/10  ']),
                         ['TYPE : 7P.279      6/10'])


class TestCleLibelle(unittest.TestCase):

    def test_formes_du_plan_et_du_pet(self):
        for libelle in ('N° PLAN', 'N°PLAN', 'NO PLAN', 'N°   PLAN'):
            self.assertEqual(cle_libelle(libelle), 'NO_PLAN', libelle)
        self.assertEqual(cle_libelle('P.E.T.'), 'PET')
        self.assertEqual(cle_libelle('TYPE'), 'TYPE')


class TestMotsDecor(unittest.TestCase):

    def test_mot_fixe_hors_libelles(self):
        self.assertEqual(mots_decor(['P.E.T.   :     {PET:<50}JARRETIERAGE',
                                     'NO PLAN : {NO_PLAN} | INDICE : {INDICE}']),
                         {'JARRETIERAGE'})

    def test_format_sans_decor(self):
        self.assertEqual(mots_decor(['CABLE :     {CABLE:<40}TYPE :    {TYPE}']), set())


class TestIndicesRevisions(unittest.TestCase):

    def test_premiere_colonne_sous_l_entete(self):
        lignes = ['INDICE     MODIFICATIONS     DATE', '  00    CREATION DU DOCUMENT',
                  '   A    Ajout', '  R10   PASSAGE EN RECOLEMENT', '  TP2   Rénovation',
                  '        suite de la ligne précédente', '  Dessiné par']
        self.assertEqual(indices_revisions(lignes), ['00', 'A', 'R10', 'TP2'])

    def test_sans_tableau_des_revisions(self):
        self.assertEqual(indices_revisions(['R1 hors tableau']), [])


# ── Pages vectorielles ────────────────────────────────────────────────

class TestPE137PiedsIdentiquesALOriginal(unittest.TestCase):
    """Les 48 pieds de 223400PE137 contre le texte du PDF, lu sans la grille."""

    @classmethod
    def setUpClass(cls):
        tpl = _tpl('REPARTITEUR')
        lecteur = pe.PdfTableExtractor(tpl)
        with fitz.open(str(PE137)) as doc:
            cls.resultats = [lecteur.extract_page_grille(doc[i], i) for i in range(len(doc))]
            cls.textes = [doc[i].get_text() for i in range(len(doc))]
        cls.pages = [(r, t) for r, t in zip(cls.resultats, cls.textes) if r['success']]
        cls.messages, cls.cellules = _excel(cls.resultats, tpl)

    def test_48_pages_indice_et_page_du_pdf(self):
        self.assertEqual(len(self.pages), 48)
        for resultat, texte in self.pages:
            indice = re.search(r'INDICE\s*:\s*(\S+)', texte).group(1)
            page = re.search(r'PAGE\s*:?\s*(\d+)', texte).group(1)
            self.assertEqual(resultat['metadata']['INDICE'], indice)
            self.assertEqual(resultat['metadata']['PAGE'], page)

    def test_indices_du_document(self):
        self.assertEqual({r['metadata']['INDICE'] for r, _ in self.pages},
                         {'10', 'R', 'R2', 'R3', 'R4', 'R9', 'R10', 'TP2'})

    def test_pied_de_l_excel_recopie_du_pdf(self):
        pieds = [c for c in self.cellules if 'INDICE' in c]
        self.assertEqual(len(pieds), 48)
        for (resultat, _), pied in zip(self.pages, pieds):
            self.assertEqual(pied, resultat['metadata']['PIED_BRUT'][-1])

    def test_bornier_absent_une_ligne_et_non_48(self):
        lignes = [m for m in self.messages if 'BORNIER' in m]
        self.assertEqual([m.strip() for m in lignes],
                         ["⚠ BORNIER absent de tout le document : vérifier le modèle"])

    def test_indices_trouves_dans_les_revisions_des_gardes(self):
        self.assertFalse([m for m in self.messages if 'INDICE' in m and '⚠' in m])

    def test_revisions_lues_sur_les_gardes_vectorielles(self):
        revisions = set().union(*(r.get('metadata', {}).get('REVISIONS', [])
                                  for r in self.resultats))
        self.assertTrue({'10', 'R', 'R2', 'R3', 'R4', 'R9', 'R10', 'TP2'} <= revisions)


class TestExtrait122a(unittest.TestCase):
    """Page 122a de 223111PE011 (vectorielle) : TYPE 7P.279 et COMPLEMENT 6/10."""

    def test_type_et_complement(self):
        with fitz.open(str(EXTRAIT)) as doc:
            meta = pe.PdfTableExtractor(_tpl()).extract_page_grille(doc[8], 8)['metadata']
        self.assertEqual((meta['TYPE'], meta['COMPLEMENT']), ('7P.279', '6/10'))

    def test_aucun_pied_faux_sur_122a_apres_excel(self):
        tpl = _tpl()
        with fitz.open(str(EXTRAIT)) as doc:
            resultat = pe.PdfTableExtractor(tpl).extract_page_grille(doc[8], 8)
        with tempfile.TemporaryDirectory() as tmp:
            sortie = Path(tmp) / 's.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                generer_excel([resultat], pe.PdfTableExtractor(tpl), sortie, on_log=print)
            converti = lire_xlsx(sortie)
        reference, colonnes = lire_verite_excel(VERITE)
        rapport = mesurer([reference[8]], converti, colonnes)
        self.assertEqual(rapport.ecarts_pieds, [])


# ── Pages lues par Claude ─────────────────────────────────────────────

REPONSE_PAGE_2 = """TYPE_PAGE: listing
G | PH QTEL2 09 | TEL PMS Q1 | PH ACC/A 01
META: {"PAGE": "2", "NO_PLAN": "223111PE011", "INDICE": "R", "CABLE": "ACC/PH02", "TYPE": "3PC200"}
PIED_BRUT: SIEMENS
PIED_BRUT: CABLE : ACC/PH02
PIED_BRUT: TYPE : 3PC200     REF CE 8707905
PIED_BRUT: N° PLAN : 223111PE011  | INDICE : R  | PAGE : 2
"""


class TestPiedBrutClaude(unittest.TestCase):

    def test_prompt_demande_le_pied_brut_et_les_revisions(self):
        prompt = _build_prompt(_tpl())
        self.assertIn('PIED_BRUT:', prompt)
        self.assertIn('REVISIONS:', prompt)

    def test_complement_retrouve(self):
        _, meta, _ = _parse_pipe_response(REPONSE_PAGE_2, _tpl())
        self.assertEqual(meta['COMPLEMENT'], 'REF CE 8707905')
        self.assertEqual(meta['TYPE'], '3PC200')
        self.assertNotIn('ALERTES_PIED', meta)

    def test_complement_dans_l_excel(self):
        _, meta, _ = _parse_pipe_response(REPONSE_PAGE_2, _tpl())
        _, cellules = _excel([_page(meta, image='page_002.png')], _tpl())
        self.assertTrue(any('REF CE 8707905' in c for c in cellules))

    def test_desaccord_meta_et_pied_brut_signale(self):
        reponse = REPONSE_PAGE_2.replace('"INDICE": "R"', '"INDICE": "R2"')
        _, meta, _ = _parse_pipe_response(reponse, _tpl())
        self.assertEqual(meta['INDICE'], 'R')
        self.assertIn('INDICE', meta['ALERTES_PIED'][0])

    def test_sans_pied_brut_meta_json_inchangee(self):
        reponse = REPONSE_PAGE_2.split('PIED_BRUT')[0]
        _, meta, _ = _parse_pipe_response(reponse, _tpl())
        self.assertEqual(meta['TYPE'], '3PC200')
        self.assertNotIn('PIED_BRUT', meta)

    def test_revisions_d_une_page_de_garde(self):
        _, meta, genre = _parse_pipe_response(
            'TYPE_PAGE: non-listing\nMETA: {}\nREVISIONS: 00 A R TP1', _tpl())
        self.assertEqual((genre, meta['REVISIONS']), ('non-listing', ['00', 'A', 'R', 'TP1']))

    def test_rejeu_garde_les_revisions_de_la_garde(self):
        with tempfile.TemporaryDirectory() as tmp:
            journal = Path(tmp) / 'j.jsonl'
            journal.write_text(json.dumps({
                'image': 'page_001.png', 'success': False, 'error': 'non-listing',
                'metadata': {'REVISIONS': ['00', 'R']},
            }) + '\n', 'utf-8')
            resultat = LogReplayer(journal, _tpl()).replay_all()[0]
        self.assertEqual(resultat['metadata']['REVISIONS'], ['00', 'R'])


# ── Aucune valeur inventée, contrôles en alerte ───────────────────────

class TestAucuneValeurInventee(unittest.TestCase):

    def test_bornier_absent_vide_et_alerte(self):
        tpl = _tpl('REPARTITEUR')
        meta = {'PAGE': '57', 'NO_PLAN': '223400PE137', 'INDICE': 'R'}
        messages, cellules = _excel([_page(meta, image='bornier_57.png', tpl=tpl)], tpl)
        self.assertFalse(any('bornier_57' in c for c in cellules))
        self.assertFalse(any('EPEULE' in c for c in cellules))
        self.assertTrue(any('BORNIER absent de tout le document' in m for m in messages))
        self.assertTrue(any('PET absent de tout le document' in m for m in messages))

    def test_indice_absent_non_invente(self):
        tpl = _tpl()
        _, cellules = _excel([_page({'PAGE': '4', 'CABLE': 'X', 'TYPE': 'Y'})], tpl)
        pied = [c for c in cellules if 'INDICE' in c][0]
        self.assertNotRegex(pied, r'INDICE\s*:\s*0')

    def test_champ_inconnu_conserve_dans_l_excel(self):
        tpl = _tpl()
        brut = ['CABLE : X      TYPE : Y', 'ARMOIRE : AR 12',
                'N° PLAN : P  | INDICE : R | PAGE : 4']
        meta = dict(analyser_pied(brut), PIED_BRUT=brut)
        self.assertEqual(meta['ARMOIRE'], 'AR 12')
        _, cellules = _excel([_page(meta)], tpl)
        self.assertTrue(any('ARMOIRE : AR 12' in c for c in cellules))


class TestAlertesPied(unittest.TestCase):

    def _meta(self, page, indice='R', plan='223111PE011'):
        return _page({'PAGE': page, 'INDICE': indice, 'NO_PLAN': plan})

    def test_indice_absent_des_revisions_alerte_sans_modification(self):
        garde = {'success': False, 'metadata': {'REVISIONS': ['00', 'A', 'R']}}
        page = self._meta('3', indice='R7')
        alertes = alertes_pied([garde, page], [page], [])
        self.assertEqual(alertes, ["page 3 : INDICE R7 absent des révisions de la page de "
                                   "garde (non modifié)"])
        self.assertEqual(page['metadata']['INDICE'], 'R7')

    def test_sans_page_de_garde_controle_declare_impossible(self):
        page = self._meta('3')
        self.assertIn('contrôle INDICE impossible', alertes_pied([page], [page], [])[0])

    def test_plan_minoritaire_signale_non_corrige(self):
        garde = {'success': False, 'metadata': {'REVISIONS': ['R']}}
        pages = [self._meta('1'), self._meta('2'), self._meta('3', plan='223111PE012')]
        alertes = alertes_pied([garde] + pages, pages, [])
        self.assertEqual(len(alertes), 1)
        self.assertIn('page 3 : N° PLAN 223111PE012 différent de la majorité', alertes[0])
        self.assertEqual(pages[2]['metadata']['NO_PLAN'], '223111PE012')

    def test_champ_absent_de_toutes_les_pages_une_seule_ligne(self):
        garde = {'success': False, 'metadata': {'REVISIONS': ['R']}}
        pages = [self._meta('1'), self._meta('2'), self._meta('3')]
        self.assertEqual(alertes_pied([garde] + pages, pages, ['CABLE']),
                         ["CABLE absent de tout le document : vérifier le modèle"])

    def test_champ_absent_d_une_page_sur_trois_alerte_sur_cette_page(self):
        garde = {'success': False, 'metadata': {'REVISIONS': ['R']}}
        pages = [self._meta('1'), self._meta('2'), self._meta('3')]
        for p in (pages[0], pages[2]):
            p['metadata']['BORNIER'] = 'B702A'
        self.assertEqual(alertes_pied([garde] + pages, pages, ['BORNIER']),
                         ["champ BORNIER absent du pied, laissé vide : page(s) 2"])

    def test_alerte_du_lecteur_rapportee_avec_sa_page(self):
        garde = {'success': False, 'metadata': {'REVISIONS': ['R']}}
        page = _page({'PAGE': '5', 'INDICE': 'R', 'ALERTES_PIED': ['INDICE : désaccord']})
        self.assertEqual(alertes_pied([garde, page], [page], []),
                         ["page 5 : INDICE : désaccord"])


class TestRenduPiedBrut(unittest.TestCase):

    def test_lignes_brutes_dans_les_deux_lignes_du_pied(self):
        tpl = _tpl()
        meta = {'PIED_BRUT': ['CABLE: WPHR/TELMG', 'TYPE : 7P.279      6/10',
                              'N°PLAN : 223111PE011     INDICE :TP3  PAGE :122a']}
        self.assertEqual(tpl.render_footer_row1(meta),
                         'CABLE: WPHR/TELMG     TYPE : 7P.279      6/10')
        self.assertEqual(tpl.render_footer_row2(meta),
                         'N°PLAN : 223111PE011     INDICE :TP3  PAGE :122a')

    def test_sans_pied_brut_format_du_modele(self):
        tpl = _tpl()
        self.assertTrue(tpl.render_footer_row1({'CABLE': 'X', 'TYPE': 'Y'}).startswith('CABLE'))


if __name__ == '__main__':
    unittest.main()


# ── Logo du bloc gauche : texte du document, jamais le libellé du modèle ──

class TestLogoPied(unittest.TestCase):

    def test_formes_du_logo(self):
        from pied_page import logo_pied
        cas = {
            'lettres espacées': ['|  |  P.E.T. : GRAND-BUT  |', '|   M A T R A   |------|'],
            'mot seul': ['PET :GRAND-BUT', 'MATRA', 'N° PLAN : 223400 PE 137'],
            'vertical avec cadre': ['M | CABLE : ACC/PH01', 'A | TYPE : 2P.279', 'T |', 'R |',
                                    'A | N° PLAN : 223111PE011'],
            'vertical sans cadre': ['M   CABLE : ACC/PH01', 'A   TYPE : 2P.279', 'T', 'R',
                                    'A   N° PLAN : 223111PE011'],
        }
        for nom, lignes in cas.items():
            self.assertEqual(logo_pied(lignes, 'M  T  I'), 'M A T R A', nom)

    def test_sans_logo_vide(self):
        from pied_page import logo_pied
        self.assertEqual(logo_pied(['CABLE : X', 'N° PLAN : P  PAGE : 3'], 'M  T  I'), '')

    def test_lettres_du_logo_hors_complement(self):
        lignes = ['M   CABLE : ACC/PH01', 'A   TYPE : 2P.279   8/10', 'T', 'R',
                  'A   N° PLAN : 223111PE011   INDICE : R   PAGE : 1']
        meta = analyser_pied(nettoyer_lignes_pied(lignes, 'SIEMENS'))
        self.assertEqual(meta['COMPLEMENT'], '8/10')


class TestLogoDansLExcel(unittest.TestCase):

    def test_pe137_matra_sur_les_48_pages(self):
        tpl = _tpl('REPARTITEUR')
        lecteur = pe.PdfTableExtractor(tpl)
        with fitz.open(str(PE137)) as doc:
            resultats = [lecteur.extract_page_grille(doc[i], i) for i in range(len(doc))]
        self.assertEqual(sum(1 for r in resultats if r['success']), 48)
        with tempfile.TemporaryDirectory() as tmp:
            sortie = Path(tmp) / 's.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                generer_excel(resultats, lecteur, sortie, on_log=print)
            ws = openpyxl.load_workbook(sortie).worksheets[0]
            gauche = [ws.cell(row=c.row - 1, column=1).value for c in ws['B']
                      if isinstance(c.value, str) and 'INDICE' in c.value]
        self.assertEqual(gauche, ['M A T R A'] * 48)

    def test_page_122a_texte_du_document(self):
        with fitz.open(str(EXTRAIT)) as doc:
            meta = pe.PdfTableExtractor(_tpl()).extract_page_grille(doc[8], 8)['metadata']
        self.assertEqual(meta['LOGO'].replace(' ', ''), 'SIEMENS')

    def test_sans_logo_cellule_vide_et_pas_le_libelle_du_modele(self):
        tpl = _tpl('REPARTITEUR')
        _, cellules = _excel([_page({'PAGE': '4', 'INDICE': 'R'}, tpl=tpl)], tpl)
        self.assertNotIn(tpl.footer_left_label, cellules)

    def test_claude_ligne_logo_et_logo_vertical(self):
        vertical = REPONSE_PAGE_2.replace('PIED_BRUT: SIEMENS\n', '').replace(
            'PIED_BRUT: CABLE', 'PIED_BRUT: M | CABLE')
        _, meta, _ = _parse_pipe_response(vertical + 'LOGO: MATRA\n', _tpl())
        self.assertEqual(meta['LOGO'], 'M A T R A')
        _, meta, _ = _parse_pipe_response(REPONSE_PAGE_2 + 'LOGO: M\n', _tpl())
        self.assertIn('LOGO', meta['ALERTES_PIED'][0])

    def test_prompt_demande_le_logo(self):
        self.assertIn('LOGO:', _build_prompt(_tpl()))


class TestTesseractIndiceNonInvente(unittest.TestCase):

    def test_indice_absent_reste_absent(self):
        from ocr_processor import BornierTableExtractor
        from template import DEFAULT_TEMPLATE
        ex = BornierTableExtractor(template=DEFAULT_TEMPLATE)
        self.assertNotIn('INDICE', ex._extract_meta([[{'text': 'PAGE : 12'}]]))
        self.assertEqual(ex._extract_meta([[{'text': 'INDICE : O | PAGE : 12'}]])['INDICE'],
                         '0')
