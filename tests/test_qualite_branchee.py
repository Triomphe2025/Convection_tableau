"""
Contrôle qualité branché (étape 9, après la mesure C0) : règles cohérence + lexique, colonne
obligatoire vide et format actives ; continuité et dictionnaire coupés. Une alerte par
cellule (toutes les raisons réunies), proposition du lexique sans remplacement, lexique
absent sans plantage, Excel sans correction du dictionnaire, validation après les contrôles
(pages à alertes d'abord, triées par nombre d'alertes).

Lancement :
    env\\Scripts\\python.exe -m pytest tests/test_qualite_branchee.py -v
"""
import contextlib
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from config import Config
from converter import Converter
from generer_classeur import (alertes_par_cellule, entrees_cellules,
                              entrees_syntheses_qualite, generer_excel)
from pdf_extractor import PdfTableExtractor
from template import TemplateManager

TPL = TemplateManager().get('REPARTITEUR 2')   # FIL, TENANT, SIGNAL, ABOUTISSANT


def _page(numero, lignes, methode='claude-vision'):
    return {'success': True, 'headers': list(TPL.columns), 'detection_method': methode,
            'metadata': {'PAGE': numero},
            'rows': [{'type': 'data', 'cells': list(c), 'confidence': [100] * 4}
                     for c in lignes]}


def _document():
    """3 fois « PH A104 01 », une fois « PH A1O4 01 » (O/0) en p. 4 ligne 1."""
    return [_page('1', [('1', 'PH A104 01', 'SIG', 'RA B1 17')]),
            _page('2', [('2', 'PH A104 01', 'SIG', 'RA B1 18')]),
            _page('3', [('3', 'PH A104 01', 'SIG', 'RA B1 19')]),
            _page('4', [('4', 'PH A1O4 01', 'SIG', 'RA B1 20')])]


def _converter(journal, dossier=None):
    conv = Converter(word_file=Path('x.pdf'), output_dir=Path(tempfile.gettempdir()),
                     template=TPL, on_log=journal.append)
    conv._lexique = set() if dossier is None else None
    # La validation enrichit le dictionnaire réel (data_dictionary.json) : jamais en test.
    conv._apprendre_corrections = lambda result: None
    return conv


class TestReglesBranchees(unittest.TestCase):

    def test_regles_actives_apres_c0(self):
        self.assertEqual(Config.QUALITE_REGLES, {'coherence': True, 'dictionnaire': False,
                                                 'continuite': False, 'colonne_vide': True,
                                                 'format': True})

    def test_dictionnaire_coupe_dans_tous_les_modes(self):
        self.assertFalse(any(Config.QUALITE_DICTIONNAIRE_MODES.values()))


class TestControlerQualite(unittest.TestCase):

    def test_alerte_posee_sur_la_cellule_avec_proposition(self):
        journal, pages = [], _document()
        _converter(journal)._controler_qualite(pages)
        self.assertEqual(pages[3]['a_verifier'], [{
            'ligne': 0, 'colonne': 1, 'type': 'cohérence',
            'message': 'lu « PH A1O4 01 », peut-être « PH A104 01 » (écrit ainsi 3 fois dans '
                       'le document)', 'lecture': None, 'qualite': True,
            'proposition': 'PH A104 01'}])
        self.assertEqual(pages[3]['rows'][0]['cells'][1], 'PH A1O4 01')
        self.assertIn('  ⚠ p. 4 ligne 1, TENANT : lu « PH A1O4 01 », peut-être « PH A104 01 » '
                      '(écrit ainsi 3 fois dans le document)', journal)

    def test_relance_sans_doublon_alertes_de_conservation_gardees(self):
        pages = _document()
        omission = {'ligne': 0, 'colonne': 3, 'type': 'élément peut-être omis',
                    'message': 'élément peut-être omis', 'lecture': 'B'}
        pages[3]['a_verifier'] = [omission]
        conv = _converter([])
        conv._controler_qualite(pages)
        conv._controler_qualite(pages)
        self.assertEqual([a['type'] for a in pages[3]['a_verifier']],
                         ['élément peut-être omis', 'cohérence'])

    def test_proposition_du_lexique(self):
        pages = [_page('119', [('1', 'CC FS 14 50', 'SIG', 'RA B1 17')])]
        conv = _converter([])
        conv._lexique = {'OC FS 14 50'}
        conv._controler_qualite(pages)
        self.assertIn('lu « CC FS 14 50 », peut-être « OC FS 14 50 » (lexique de référence)',
                      pages[0]['a_verifier'][0]['message'])
        self.assertEqual(pages[0]['rows'][0]['cells'][1], 'CC FS 14 50')

    def test_lexique_absent_une_ligne_au_journal_pas_de_plantage(self):
        journal = []
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(Config, 'QUALITE_LEXIQUE_DOSSIER', Path(tmp) / 'absent'):
            conv = _converter(journal, dossier=True)
            conv._controler_qualite(_document())
        self.assertEqual(conv._lexique, set())
        self.assertEqual(sum('lexique de référence absent ou vide' in m for m in journal), 1)

    def test_lexique_illisible_cohérence_seule(self):
        journal = []
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / 'abime.xlsx').write_bytes(b'pas un classeur')
            with patch.object(Config, 'QUALITE_LEXIQUE_DOSSIER', Path(tmp)):
                pages = _document()
                _converter(journal, dossier=True)._controler_qualite(pages)
        self.assertTrue(any('lexique de référence illisible' in m for m in journal))
        self.assertEqual(len(pages[3]['a_verifier']), 1)

    def test_continuite_non_branchee(self):
        pages = [_page('1', [('1', 'PH A104 01', 'SIG', 'RA B1 17'),
                             ('2', 'PH A104 01', 'SIG', 'RA B1 90')])]
        _converter([])._controler_qualite(pages)
        self.assertFalse(pages[0].get('a_verifier'))

    def test_synthese_du_garde_fou_au_journal_et_a_part(self):
        journal = []
        pages = [_page(str(n), [(str(n), 'texte libre', 'SIG', 'RA B1 17')]) for n in (1, 2)]
        _converter(journal)._controler_qualite(pages)
        self.assertEqual(pages[0]['syntheses_qualite'], [{
            'type': 'format', 'message': 'format de la colonne TENANT non reconnu dans ce '
                                         'document, contrôle coupé pour cette colonne'}])
        self.assertIn('  ⚠ contrôle qualité : format de la colonne TENANT non reconnu dans ce '
                      'document, contrôle coupé pour cette colonne', journal)


class TestUneAlerteParCellule(unittest.TestCase):

    def _resultat(self):
        return {'headers': list(TPL.columns), 'metadata': {'PAGE': '7'}, 'a_verifier': [
            {'ligne': 0, 'colonne': 1, 'type': 'élément peut-être omis',
             'message': 'élément peut-être omis : Tesseract lit « 02 »', 'lecture': '02'},
            {'ligne': 0, 'colonne': 1, 'type': 'cohérence, format', 'qualite': True,
             'message': 'lu « X », peut-être « Y » ; « X » hors du format de TENANT',
             'lecture': None},
            {'ligne': 0, 'colonne': 0, 'type': 'ligne peut-être manquante',
             'message': 'ligne peut-être manquante', 'lecture': 'A0'},
            {'ligne': None, 'colonne': 0, 'type': 'colonne vide', 'qualite': True,
             'message': 'colonne FIL entièrement vide sur la page', 'lecture': None}]}

    def test_une_ligne_qui_donne_toutes_les_raisons(self):
        alertes = alertes_par_cellule(self._resultat())
        self.assertEqual([a['type'] for a in alertes],
                         ['élément peut-être omis, cohérence, format',
                          'ligne peut-être manquante', 'colonne vide'])
        self.assertEqual(alertes[0]['message'], 'élément peut-être omis : Tesseract lit « 02 » '
                         '; lu « X », peut-être « Y » ; « X » hors du format de TENANT')
        self.assertEqual(alertes[0]['lecture'], '02')

    def test_alerte_de_page_sans_numero_de_ligne(self):
        entree = entrees_cellules([self._resultat()])[-1]
        self.assertEqual((entree['page'], entree['ligne'], entree['colonne'], entree['cible']),
                         ('7', None, 'FIL', (0, None, 0)))

    def test_syntheses_sans_cible(self):
        self.assertEqual(entrees_syntheses_qualite([{'syntheses_qualite': [
            {'type': 'format', 'message': 'm'}]}, {}]), [
            {'page': None, 'ligne': None, 'colonne': None, 'type': 'format', 'message': 'm',
             'lecture': None, 'cible': None}])


class TestExcel(unittest.TestCase):

    def _ecrire(self, pages, appels=None):
        original = PdfTableExtractor._fill_worksheet

        def espion(extracteur, *args, **kw):
            if appels is not None:
                appels.append(kw.get('dictionary'))
            return original(extracteur, *args, **kw)

        tmp = tempfile.mkdtemp()
        sortie = Path(tmp) / 's.xlsx'
        with contextlib.redirect_stdout(io.StringIO()), \
                patch.object(PdfTableExtractor, '_fill_worksheet', espion):
            generer_excel(pages, PdfTableExtractor(TPL), sortie, on_log=lambda m: None)
        return openpyxl.load_workbook(sortie)

    def test_dictionnaire_jamais_passe_au_remplissage(self):
        appels = []
        self._ecrire([_page('1', [('1', 'PH A104 01', 'SIG', 'RA B1 17')], methode='header')],
                     appels)
        self.assertEqual(appels, [None])

    def test_cellule_qualite_coloree_commentee_valeur_gardee(self):
        pages = _document()
        _converter([])._controler_qualite(pages)
        wb = self._ecrire(pages)
        cellule = next(c for ligne in wb['Borniers'].iter_rows() for c in ligne
                       if 'A1O4' in str(c.value))
        self.assertEqual(cellule.column_letter, 'B')
        self.assertEqual(' '.join(cellule.value.split()), 'PH A1O4 01')
        self.assertEqual(cellule.fill.fgColor.rgb, '00' + Config.COULEUR_QUALITE)
        self.assertIn('peut-être « PH A104 01 »', cellule.comment.text)
        lignes = [[c.value for c in r] for r in wb['A VERIFIER'].iter_rows()]
        self.assertEqual(lignes[1][:4], ['4', 1, 'TENANT', 'cohérence'])

    def test_cellule_corrigee_b3_garde_son_orange_une_seule_ligne(self):
        tpl = TemplateManager().get('REPARTITEUR')   # TENANT, JAR, ABOUTISSANT, SIGNAL
        page = {'success': True, 'headers': list(tpl.columns), 'metadata': {'PAGE': '18'},
                'detection_method': 'pdf-grille',
                'rows': [{'type': 'data', 'cells': [t, '0815B', 'D3T 02A', 'SIG'],
                          'confidence': [100] * 4}
                         for t in ('D3T O1A', 'D3T 03A', 'D3T 04A', 'D3T 05A', 'D3T 06A',
                                   'D3T 07A')]}
        conv = Converter(word_file=Path('x.pdf'), output_dir=Path(tempfile.gettempdir()),
                         template=tpl, on_log=lambda m: None)
        conv._lexique = set()
        conv._controler_qualite([page])
        sortie = Path(tempfile.mkdtemp()) / 's.xlsx'
        with contextlib.redirect_stdout(io.StringIO()):
            generer_excel([page], PdfTableExtractor(tpl), sortie, on_log=lambda m: None)
        wb = openpyxl.load_workbook(sortie)
        cellule = next(c for ligne in wb['Borniers'].iter_rows() for c in ligne
                       if 'D3T' in str(c.value) and c.column == 1)
        self.assertEqual(' '.join(cellule.value.split()), 'D3T 01A')
        self.assertEqual(cellule.fill.fgColor.rgb, '00' + Config.COULEUR_CORRIGE)
        lignes = [[c.value for c in r] for r in wb['A VERIFIER'].iter_rows()][1:]
        self.assertEqual([(r[3], r[4]) for r in lignes], [(
            'format', "« D3T O1A » hors du format de TENANT ; corrigé dans l'Excel "
                      "(coquille O/0), l'original porte O1A")])

    def test_synthese_en_tete_de_a_verifier(self):
        pages = [_page(str(n), [(str(n), 'texte libre', 'SIG', 'RA B1 17')]) for n in (1, 2)]
        _converter([])._controler_qualite(pages)
        lignes = [[c.value for c in r] for r in self._ecrire(pages)['A VERIFIER'].iter_rows()]
        self.assertEqual(lignes[1][:5], [None, None, None, 'format',
                                         'format de la colonne TENANT non reconnu dans ce '
                                         'document, contrôle coupé pour cette colonne'])


class TestValidationApresControles(unittest.TestCase):

    def _inscrire(self, conv, pages, relire=None):
        conv._a_valider = [{'result': r, 'image': lambda: 'page.png', 'relire': relire,
                            'moteur': 'Claude Vision', 'page_pdf': None} for r in pages]

    def test_pages_a_alertes_d_abord_triees_par_nombre(self):
        vus = []
        pages = [_page(str(n), [(str(n), 'PH A104 01', 'SIG', 'RA B1 17')]) for n in (1, 2, 3, 4)]
        pages[1]['a_verifier'] = [{'ligne': 0}]
        pages[3]['a_verifier'] = [{'ligne': 0}, {'ligne': 0}]
        conv = _converter([])
        conv._on_validation = lambda rang, total, r, image: vus.append(
            (rang, total, r['metadata']['PAGE'])) or None
        self._inscrire(conv, pages)
        conv._valider_document(pages)
        self.assertEqual(vus, [(1, 4, '4'), (2, 4, '2'), (3, 4, '1'), (4, 4, '3')])

    def test_relecture_puis_nouveau_controle(self):
        pages = _document()
        conv = _converter([])
        conv._controler_qualite(pages)
        reponses = [('retry', 'le O est un 0'), None]
        vus = []

        def valider(rang, total, r, image):
            vus.append((r['_attempt'], len(r.get('a_verifier') or [])))
            return reponses.pop(0)

        conv._on_validation = valider
        self._inscrire(conv, pages, relire=lambda image, fb: _page(
            '4', [('4', 'PH A104 01', 'SIG', 'RA B1 20')]))
        conv._a_valider = conv._a_valider[3:]
        conv._valider_document(pages)
        self.assertEqual(vus, [(1, 1), (2, 0)])
        self.assertEqual(pages[3]['rows'][0]['cells'][1], 'PH A104 01')

    def test_page_en_couche_texte_sans_relecture(self):
        pages = [_page('1', [('1', 'PH A104 01', 'SIG', 'RA B1 17')], methode='pdf-grille')]
        images = []
        conv = _converter([])
        conv._on_validation = lambda rang, total, r, image: images.append(image) or (
            ('retry', 'x') if len(images) == 1 else None)
        self._inscrire(conv, pages)
        conv._valider_document(pages)
        self.assertEqual(images, [''])

    def test_sans_validation_interactive_rien(self):
        conv = _converter([])
        self._inscrire(conv, _document())
        conv._valider_document(_document())


if __name__ == '__main__':
    unittest.main()
