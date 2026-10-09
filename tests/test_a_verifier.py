"""
Feuille « A VERIFIER » de l'Excel livré (commit B1, conçue pour l'étape 9) : colonnes page,
ligne, colonne, type, message, lecture Tesseract, lien cliquable vers la cellule ou le bloc
de la page. Pages dégradées : une ligne par page, ou une synthèse suivie des pages triées par
taux décroissant quand plus de la moitié des pages scannées sont dégradées.

Lancement :
    env\\Scripts\\python.exe -m pytest tests/test_a_verifier.py -v
"""
import contextlib
import io
import tempfile
import unittest
from pathlib import Path

import openpyxl

from config import Config
from generer_classeur import COLONNES_A_VERIFIER, entrees_pages_degradees, generer_excel
from pdf_extractor import PdfTableExtractor
from template import TemplateManager

TPL = TemplateManager().get('REPARTITEUR 2')


def _page(numero, taux=None):
    page = {'success': True, 'headers': list(TPL.columns), 'image_path': f'page_{numero}.png',
            'detection_method': 'claude-vision', 'metadata': {'PAGE': numero},
            'rows': [{'type': 'data', 'cells': ['A1', 'PH A104 01', 'SIG', ''],
                      'confidence': [100] * 4}]}
    if taux is not None:
        page['controle_page'] = {'taux': taux, 'degradee': taux > Config.PAGE_DEGRADEE}
    return page


def _feuille(pages):
    with tempfile.TemporaryDirectory() as tmp:
        sortie = Path(tmp) / 's.xlsx'
        with contextlib.redirect_stdout(io.StringIO()):
            generer_excel(pages, PdfTableExtractor(TPL), sortie, on_log=lambda m: None)
        wb = openpyxl.load_workbook(sortie)
        if 'A VERIFIER' not in wb.sheetnames:
            return wb.sheetnames, None
        ws = wb['A VERIFIER']
        lignes = [[c.value for c in r] for r in ws.iter_rows()]
        liens = [r[-1].hyperlink.location if r[-1].hyperlink else None for r in ws.iter_rows()]
        return wb.sheetnames, (lignes, liens)


class TestEntreesPagesDegradees(unittest.TestCase):

    def test_minorite_une_ligne_par_page(self):
        entrees = entrees_pages_degradees([_page('1', 0.02), _page('52', 0.29),
                                           _page('3', 0.05), _page('4', 0.01)])
        self.assertEqual([(e['page'], e['type'], e['message']) for e in entrees],
                         [('52', 'page dégradée',
                           'scan dégradé (divergence 29 %) : à relire en priorité')])
        self.assertEqual(entrees[0]['cible'], (1, None, None))

    def test_majorite_une_synthese_puis_pages_triees(self):
        entrees = entrees_pages_degradees([_page('1', 0.18), _page('2', 0.50),
                                           _page('3', 0.29), _page('4', 0.02)])
        self.assertEqual(entrees[0]['type'], 'synthèse')
        self.assertEqual(entrees[0]['message'], 'scan dégradé sur 3 pages sur 4 (taux de 18 à '
                                                '50 %) : contrôle de conservation impossible')
        self.assertEqual([(e['page'], e['message']) for e in entrees[1:]],
                         [('2', 'divergence 50 %'), ('3', 'divergence 29 %'),
                          ('1', 'divergence 18 %')])

    def test_moitie_exacte_pas_de_synthese(self):
        entrees = entrees_pages_degradees([_page('1', 0.18), _page('2', 0.02)])
        self.assertEqual([e['type'] for e in entrees], ['page dégradée'])

    def test_pages_sans_mesure_hors_du_compte(self):
        # Pages vectorielles (pas de lecture Tesseract) : ni scannées ni dégradées.
        entrees = entrees_pages_degradees([_page('1', 0.18), _page('2'), _page('3'),
                                           _page('4', 0.02)])
        self.assertEqual([e['type'] for e in entrees], ['page dégradée'])
        entrees = entrees_pages_degradees([_page('1', 0.18), _page('2', 0.30), _page('3'),
                                           _page('4', 0.02)])
        self.assertTrue(entrees[0]['message'].startswith('scan dégradé sur 2 pages sur 3 '))

    def test_aucune_page_degradee(self):
        self.assertEqual(entrees_pages_degradees([_page('1', 0.02), _page('2')]), [])


class TestFeuilleAVerifier(unittest.TestCase):

    def test_colonnes_conçues_pour_l_etape_9(self):
        self.assertEqual(COLONNES_A_VERIFIER, ['Page', 'Ligne', 'Colonne', 'Type', 'Message',
                                               'Lecture Tesseract', 'Lien'])

    def test_ligne_et_lien_vers_le_bloc_de_la_page(self):
        noms, (lignes, liens) = _feuille([_page('1', 0.02), _page('52', 0.29)])
        self.assertEqual(noms[0], 'Borniers')
        self.assertEqual(lignes[0], COLONNES_A_VERIFIER)
        self.assertEqual(lignes[1][:5], ['52', None, None, 'page dégradée',
                                         'scan dégradé (divergence 29 %) : à relire en '
                                         'priorité'])
        # Le bloc de la 2e page commence après PAGE_SIZE lignes : en-tête en A49.
        self.assertEqual(liens[1], f"'Borniers'!A{1 + Config.PAGE_SIZE}")

    def test_synthese_sans_lien_puis_pages_liees(self):
        _, (lignes, liens) = _feuille([_page('1', 0.18), _page('2', 0.50), _page('3', 0.02)])
        self.assertEqual(lignes[1][3], 'synthèse')
        self.assertIsNone(liens[1])
        self.assertEqual([r[0] for r in lignes[2:]], ['2', '1'])
        self.assertEqual(liens[2:], [f"'Borniers'!A{1 + Config.PAGE_SIZE}", "'Borniers'!A1"])

    def test_pas_de_feuille_sans_entree(self):
        noms, feuille = _feuille([_page('1', 0.02), _page('2')])
        self.assertEqual((noms, feuille), (['Borniers'], None))


if __name__ == '__main__':
    unittest.main()


class TestAlertesDeConservation(unittest.TestCase):
    """Entrées « a_verifier » d'une page (contrôle de conservation, étape 9 ensuite)."""

    def _pages(self):
        page = _page('7', 0.02)
        page['rows'].append({'type': 'data', 'cells': ['A2', 'PH A104', 'SIG', ''],
                             'confidence': [100] * 4})
        page['a_verifier'] = [
            {'ligne': 1, 'colonne': 1, 'type': 'élément peut-être omis',
             'message': 'élément peut-être omis : Tesseract lit « 02 » (confiance 91)',
             'lecture': '02'},
            {'ligne': 1, 'colonne': 0, 'type': 'ligne peut-être manquante',
             'message': 'ligne peut-être manquante entre « A1 | PH A104 01 | SIG | » et '
                        '« A2 | PH A104 | SIG | »', 'lecture': 'A15 PH A104 15'},
        ]
        return [_page('6', 0.29), page]

    def test_lignes_apres_les_pages_avec_lien_vers_la_cellule(self):
        _, (lignes, liens) = _feuille(self._pages())
        self.assertEqual([r[3] for r in lignes[1:]], ['page dégradée', 'élément peut-être omis',
                                                      'ligne peut-être manquante'])
        self.assertEqual(lignes[2][:3], ['7', 2, 'TENANT'])
        self.assertEqual(lignes[2][5], '02')
        debut = 1 + Config.PAGE_SIZE
        self.assertEqual(liens[2:], [f"'Borniers'!B{debut + 2}", f"'Borniers'!A{debut + 2}"])

    def test_cellule_coloree_commentee_valeur_de_claude_gardee(self):
        with tempfile.TemporaryDirectory() as tmp:
            sortie = Path(tmp) / 's.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                generer_excel(self._pages(), PdfTableExtractor(TPL), sortie,
                              on_log=lambda m: None)
            ws = openpyxl.load_workbook(sortie)['Borniers']
            debut = 1 + Config.PAGE_SIZE
            omise, suivante = ws[f'B{debut + 2}'], ws[f'A{debut + 2}']
            self.assertEqual(' '.join(omise.value.split()), 'PH A104')
            self.assertEqual(omise.fill.fgColor.rgb, '00' + Config.COULEUR_CONSERVATION)
            self.assertIn('« 02 »', omise.comment.text)
            self.assertIn('ligne peut-être manquante', suivante.comment.text)
            self.assertEqual(suivante.value.strip(), 'A2')

    def test_page_degradee_minoritaire_controle_impossible(self):
        page = _page('52', 0.29)
        page['controle_page']['controle_impossible'] = True
        entrees = entrees_pages_degradees([page, _page('1', 0.02), _page('2', 0.03)])
        self.assertEqual(entrees[0]['message'], 'scan dégradé (divergence 29 %) : à relire en '
                                                'priorité ; contrôle de conservation impossible')
