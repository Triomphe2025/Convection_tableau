"""
Pages écartées du classeur : raison écrite au journal, tableau vide conservé.

Lancement :
    env\\Scripts\\python.exe -m pytest tests/test_pages_ignorees.py -v
"""
import contextlib
import io
import tempfile
import unittest
from pathlib import Path

import fitz
import openpyxl

import pdf_extractor as pe
from config import Config
from generer_classeur import (_ERREUR_HORS_TABLEAU, generer_excel, numero_page,
                              raison_page_ignoree)
from template import TemplateManager

FIXTURES = Path(__file__).parent / 'fixtures'
EXTRAIT = FIXTURES / '223111PE011_extrait_10pages.pdf'
COLONNES = ['FIL', 'TENANT', 'SIGNAL', 'ABOUTISSANT']


def _page(n_lignes, page='7', cellules=('N', 'PH AD 03', 'RESERVE', 'PH AL 01')):
    return {
        'success': True, 'headers': list(COLONNES), 'metadata': {'PAGE': page},
        'image_path': f'page_{page}', 'detection_method': 'pdf-grille', 'blur_pct': 0.0,
        'rows': [{'type': 'data', 'cells': list(cellules), 'confidence': [100] * 4}
                 for _ in range(n_lignes)],
    }


def _raison(resultat, min_rows=1, densite=0.15):
    return raison_page_ignoree(resultat, COLONNES, min_rows, densite)


class TestConfig(unittest.TestCase):

    def test_min_data_rows_vaut_1(self):
        self.assertEqual(Config.MIN_DATA_ROWS, 1)


class TestNumeroPage(unittest.TestCase):

    def test_numero_du_pied(self):
        self.assertEqual(numero_page({'metadata': {'PAGE': '122a'}}), '122a')

    def test_repli_sur_l_image(self):
        self.assertEqual(numero_page({'metadata': {}, 'image_path': 'x/page_003.png'}),
                         'page_003.png')

    def test_repli_sur_la_position(self):
        self.assertEqual(numero_page({'page_num': 4}), '4')


class TestRaisonPageIgnoree(unittest.TestCase):

    def test_page_de_une_ligne_conservee(self):
        self.assertIsNone(_raison(_page(1)))

    def test_tableau_vide_conserve(self):
        self.assertIsNone(_raison(_page(0), min_rows=3))

    def test_sous_le_minimum_ecartee_avec_raison(self):
        self.assertEqual(_raison(_page(2), min_rows=3),
                         "2 ligne(s) de données, minimum 3 (MIN_DATA_ROWS)")

    def test_page_hors_tableau(self):
        raison = _raison({'success': False, 'error': _ERREUR_HORS_TABLEAU})
        self.assertIn('pas un tableau de câblage', raison)

    def test_erreur_du_moteur_recopiee(self):
        self.assertEqual(_raison({'success': False, 'error': 'réponse tronquée'}),
                         'réponse tronquée')

    def test_echec_sans_erreur(self):
        self.assertIn('lecture impossible', _raison({'success': False}))

    def test_en_tete_etranger_au_modele(self):
        resultat = dict(_page(3), headers=['A', 'B'])
        self.assertIn('en-tête', _raison(resultat))

    def test_gribouillage_ecarte(self):
        resultat = _page(3, cellules=('|', '-', '', '.'))
        self.assertIn('PAGE_DENSITE_MIN', _raison(resultat))


class TestGenererExcelJournal(unittest.TestCase):

    def _generer(self, resultats):
        messages = []
        tpl = TemplateManager().get('REPARTITEUR 2')
        with tempfile.TemporaryDirectory() as tmp:
            sortie = Path(tmp) / 'sortie.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                generer_excel(resultats, pe.PdfTableExtractor(tpl), sortie,
                              on_log=messages.append)
            valeurs = {
                str(v) for ligne in openpyxl.load_workbook(sortie).worksheets[0].iter_rows(
                    values_only=True) for v in ligne if v is not None
            } if sortie.exists() else set()
        return messages, valeurs

    def test_page_ignoree_ecrite_au_journal(self):
        messages, _ = self._generer([_page(2), {'success': False, 'page_num': 9,
                                                'error': _ERREUR_HORS_TABLEAU}])
        ignorees = [m for m in messages if 'page ignorée' in m]
        self.assertEqual(len(ignorees), 1)
        self.assertTrue(ignorees[0].strip().startswith('page ignorée : 9 pas un tableau'))

    def test_page_de_deux_lignes_dans_le_classeur(self):
        messages, valeurs = self._generer([_page(2, cellules=('B', 'PH Q 04', 'X', 'Y'))])
        self.assertIn('PH Q 04', valeurs)
        self.assertFalse([m for m in messages if 'page ignorée' in m])


class TestTableauVideReserve(unittest.TestCase):
    """Page 104 de 223111PE011 : en-tête, aucune ligne, CABLE : RESERVE."""

    def setUp(self):
        self.tpl = TemplateManager().get('REPARTITEUR 2')
        with fitz.open(str(EXTRAIT)) as doc:
            self.resultat = pe.PdfTableExtractor(self.tpl).extract_page_grille(doc[5], 5)

    def test_grille_rend_un_tableau_vide(self):
        self.assertTrue(self.resultat['success'])
        self.assertEqual([r for r in self.resultat['rows'] if r['type'] == 'data'], [])
        self.assertEqual(self.resultat['metadata']['CABLE'], 'RESERVE')

    def test_conserve_dans_le_classeur(self):
        messages = []
        with tempfile.TemporaryDirectory() as tmp:
            sortie = Path(tmp) / 'vide.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                generer_excel([self.resultat], pe.PdfTableExtractor(self.tpl), sortie,
                              on_log=messages.append)
            texte = ' '.join(
                str(v) for ligne in openpyxl.load_workbook(sortie).worksheets[0].iter_rows(
                    values_only=True) for v in ligne if v is not None)
        self.assertIn('RESERVE', texte)
        self.assertIn('104', texte)
        self.assertFalse([m for m in messages if 'page ignorée' in m])

    def test_page_sans_en_tete_a_une_raison(self):
        doc = fitz.open()
        doc.new_page().insert_text((72, 72), "FEUILLE DE MODIFICATIONS  INDICE A")
        resultat = pe.PdfTableExtractor(self.tpl).extract_page_grille(doc[0], 0)
        doc.close()
        self.assertFalse(resultat['success'])
        self.assertIn("pas d'en-tête du modèle", _raison(resultat))


if __name__ == '__main__':
    unittest.main()
