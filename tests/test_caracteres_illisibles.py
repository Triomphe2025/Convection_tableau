"""
Caractère illisible : Claude écrit « ?? », gardé tel quel dans l'Excel et coloré.

Lancement :
    env\\Scripts\\python.exe -m pytest tests/test_caracteres_illisibles.py -v
"""
import contextlib
import io
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl import Workbook

from claude_ocr import _build_prompt, _parse_pipe_response
from config import Config
from generer_classeur import generer_excel, marquer_illisibles
from ocr_processor import BornierTableExtractor
from template import DEFAULT_TEMPLATE

REPONSE = (
    "TYPE_PAGE: listing\n"
    "01 | BL?? | TC IDPO1 | 0109R\n"
    "02 | ROUGE | PH Q??L2 09 | 0110W\n"
    'META: {"PAGE": "12", "BORNIER": "PA"}\n'
)


class TestPrompt(unittest.TestCase):

    def test_consigne_exacte(self):
        self.assertIn(
            "Si un caractère est illisible, écris ?? à sa place. Ne devine jamais,"
            " ne corrige jamais un mot, recopie exactement.",
            _build_prompt(DEFAULT_TEMPLATE),
        )


class TestParseurGardeLeMarqueur(unittest.TestCase):

    def test_marqueur_intact(self):
        rows, _, _ = _parse_pipe_response(REPONSE, DEFAULT_TEMPLATE)
        self.assertEqual(rows[0]['cells'][1], 'BL??')
        self.assertEqual(rows[1]['cells'][2], 'PH Q??L2 09')


class TestMarquerIllisibles(unittest.TestCase):

    def test_seules_les_cellules_marquees_sont_colorees(self):
        ws = Workbook().active
        ws['A1'], ws['B1'], ws['C1'] = 'BL??', 'BLANC', None
        self.assertEqual(marquer_illisibles(ws, '??', 'FFC7CE'), 1)
        self.assertEqual(ws['A1'].fill.fgColor.rgb, '00FFC7CE')
        self.assertIn('illisible', ws['A1'].comment.text)
        self.assertIsNone(ws['B1'].fill.fill_type)
        self.assertIsNone(ws['B1'].comment)

    def test_aucune_cellule(self):
        ws = Workbook().active
        ws['A1'] = 'PH QTEL2 09'
        self.assertEqual(marquer_illisibles(ws, '??', 'FFC7CE'), 0)


class TestExcelDeBoutEnBout(unittest.TestCase):

    def test_marqueur_garde_colore_et_signale(self):
        rows, meta, _ = _parse_pipe_response(REPONSE, DEFAULT_TEMPLATE)
        resultat = {
            'success': True, 'headers': list(DEFAULT_TEMPLATE.columns), 'rows': rows,
            'metadata': meta, 'image_path': 'page_012.png', 'detection_method': 'claude-vision',
        }
        messages = []
        with tempfile.TemporaryDirectory() as tmp:
            sortie = Path(tmp) / 's.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                generer_excel([resultat], BornierTableExtractor(template=DEFAULT_TEMPLATE),
                              sortie, on_log=messages.append)
            ws = openpyxl.load_workbook(sortie).worksheets[0]
            marquees = {c.value: c.fill.fgColor.rgb for ligne in ws.iter_rows()
                        for c in ligne if isinstance(c.value, str) and '??' in c.value}
        self.assertEqual(set(marquees), {'BL??', 'PH Q??L2 09'})
        self.assertTrue(all(rgb.endswith(Config.COULEUR_ILLISIBLE) for rgb in marquees.values()))
        self.assertTrue(any('2 cellule(s) avec « ?? »' in m for m in messages))


if __name__ == '__main__':
    unittest.main()
