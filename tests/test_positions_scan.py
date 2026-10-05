"""
Commit A — positions d'origine des mots sur les pages scannées (positions_scan.py),
sur des lectures Tesseract fabriquées : grille de caractères (pas + phase), coupe aux
traits du cadre, jumeaux (exacts, par rang, sans jumeau), déplacement vers la colonne
voisine, lignes de section, contenu inchangé ; rendu dans l'Excel (generer_classeur).

Lancement :
    env\\Scripts\\python.exe -m pytest tests/test_positions_scan.py -v
"""
import contextlib
import io
import random
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from config import Config
from generer_classeur import generer_excel, texte_aux_positions
from pdf_extractor import PdfTableExtractor
from positions_scan import (LecturePage, ajuster_grille, colonne_caractere, couper_aux_traits,
                            dpi_de_rendu, pas_initial, placer_page, replier)
from template import TemplateManager

PAS, PHASE = 10.0, 5.0
COLONNES = ['FIL', 'TENANT', 'SIGNAL', 'ABOUTISSANT']
TRAITS = [0.0, 300.0, 600.0, 900.0, 1200.0]       # colonne k du tableau = caractères 30k à 30k+29


def _mot(col, ligne, texte, conf=90):
    """Mot posé sur la grille fabriquée : 1er caractère en colonne `col`, ligne `ligne`."""
    x0 = PHASE + PAS * col
    return (x0, 100 * ligne, x0 + PAS * len(texte), 100 * ligne + 30, texte, conf)


def _lecture(*mots):
    return LecturePage(mots=list(mots), traits=list(TRAITS), dpi=300, taille=(1200, 2000))


def _ligne(*cellules):
    return {'type': 'data', 'cells': list(cellules), 'confidence': [100] * len(cellules)}


def _mots(row):
    return sorted(m for c in row.get('cells', [row.get('text', '')]) for m in c.split())


# Ligne de référence posée sur chaque page fabriquée : fixe l'origine de chaque colonne.
REPERE_TESS = [_mot(1, 1, 'A1'), _mot(31, 1, 'PH'), _mot(37, 1, 'A104'), _mot(47, 1, '01'),
               _mot(61, 1, 'SIG'), _mot(91, 1, 'PH'), _mot(97, 1, 'TA106'), _mot(107, 1, '01')]
REPERE_CLAUDE = _ligne('A1', 'PH A104 01', 'SIG', 'PH TA106 01')


class TestPasInitial(unittest.TestCase):

    def test_mediane_largeur_par_caractere(self):
        mots = [_mot(0, 1, 'ABCD'), _mot(10, 1, 'EFGHIJ'), _mot(20, 1, 'X'), _mot(30, 1, '||')]
        self.assertAlmostEqual(pas_initial(mots), PAS)

    def test_pente_largeur_selon_le_nombre_de_caracteres(self):
        # Boîte du dernier caractère plus étroite qu'un pas : largeur = (n - 0,2) × pas ;
        # largeur / n sous-estimerait le pas, la pente le retrouve.
        mots = [(0, 0, (n - 0.2) * 28.6, 30, 'A' * n, 90) for n in (2, 3, 4, 5, 6) for _ in '123']
        self.assertAlmostEqual(pas_initial(mots), 28.6)

    def test_aucun_mot_exploitable(self):
        self.assertIsNone(pas_initial([_mot(0, 1, '|')]))


class TestAjusterGrille(unittest.TestCase):

    def _debuts(self, pas, phase, n=400, bruit=1.0):
        alea = random.Random(7)
        return [phase + pas * alea.randint(0, 90) + alea.uniform(-bruit, bruit) for _ in range(n)]

    def test_pas_et_phase_retrouves(self):
        pas, phase = ajuster_grille(self._debuts(22.37, 7.3), 22.9)
        self.assertAlmostEqual(pas, 22.37, delta=0.03)
        self.assertAlmostEqual(phase, 7.3, delta=1.0)

    def test_demi_pas_ecarte(self):
        # Un pas moitié aligne aussi tous les débuts : la recherche reste près du pas estimé.
        pas, _ = ajuster_grille(self._debuts(29.1, 2.0), 28.2)
        self.assertAlmostEqual(pas, 29.1, delta=0.05)

    def test_colonne_arrondie(self):
        self.assertEqual(colonne_caractere(PHASE + 12 * PAS + 4, PAS, PHASE), 12)
        self.assertEqual(colonne_caractere(PHASE + 12 * PAS - 4, PAS, PHASE), 12)


class TestDpiDeRendu(unittest.TestCase):

    def test_a4_a_300_dpi(self):
        self.assertEqual(dpi_de_rendu(595, 842), Config.POSITIONS_DPI)
        self.assertEqual(dpi_de_rendu(842, 595), Config.POSITIONS_DPI)

    def test_mediabox_geante_meme_nombre_de_pixels(self):
        # 223111PE012 p. 4 : 2481 × 3505 pt → environ 2480 × 3500 px, comme un A4 à 300 DPI.
        dpi = dpi_de_rendu(2481, 3505)
        self.assertEqual(dpi, 72)
        self.assertAlmostEqual(3505 * dpi / 72, 842 * 300 / 72, delta=30)


class TestCouperAuxTraits(unittest.TestCase):

    # Trait du cadre en x = 310 : dans la colonne de caractère 30 (de 305 à 315).
    def test_coupe_au_trait_lu(self):
        mot = _mot(25, 1, '0785B|P111TC')
        coupes = couper_aux_traits([mot], [0.0, 310.0, 1200.0], PAS)
        self.assertEqual([(m[4], colonne_caractere(m[0], PAS, PHASE)) for m in coupes],
                         [('0785B', 25), ('P111TC', 31)])

    def test_trait_lu_sur_deux_caracteres(self):
        # 6A23111PE133 p. 15, ligne 31 : « 0815B/|RM », le trait n'occupe qu'une colonne.
        coupes = couper_aux_traits([_mot(25, 1, '0815B/|RM')], [0.0, 310.0, 1200.0], PAS)
        self.assertEqual([(m[4], colonne_caractere(m[0], PAS, PHASE)) for m in coupes],
                         [('0815B', 25), ('RM', 31)])

    def test_trait_lu_comme_barre_oblique(self):
        coupes = couper_aux_traits([_mot(25, 1, '0816B/RM')], [0.0, 310.0, 1200.0], PAS)
        self.assertEqual([m[4] for m in coupes], ['0816B', 'RM'])

    def test_barre_oblique_d_une_valeur_gardee(self):
        coupes = couper_aux_traits([_mot(2, 1, 'WPHR/A105')], TRAITS, PAS)
        self.assertEqual([m[4] for m in coupes], ['WPHR/A105'])

    def test_coupe_au_trait_du_cadre_non_lu(self):
        mot = _mot(28, 1, 'ABCDEF')                  # x de 285 à 345, trait en 300
        coupes = couper_aux_traits([mot], [0.0, 300.0, 1200.0], PAS)
        self.assertEqual([(m[4], colonne_caractere(m[0], PAS, PHASE)) for m in coupes],
                         [('AB', 28), ('CDEF', 30)])

    def test_trait_seul_et_mot_sans_trait(self):
        coupes = couper_aux_traits([_mot(40, 1, '|'), _mot(2, 1, 'A1')], TRAITS, PAS)
        self.assertEqual([m[4] for m in coupes], ['A1'])


class TestReplier(unittest.TestCase):

    def test_o_et_i(self):
        self.assertEqual(replier('QTEL2 O1I'), 'QTEL2 011')

    def test_autres_inchanges(self):
        self.assertEqual(replier('PH 7T'), 'PH 7T')


class TestPlacerPage(unittest.TestCase):

    def _placer(self, claude, tess):
        rows, bilan = placer_page([REPERE_CLAUDE, claude], COLONNES,
                                  _lecture(*REPERE_TESS, *tess))
        return rows[1], bilan

    def test_jumeaux_exacts(self):
        row, bilan = self._placer(_ligne('A2', 'PH A104 02', 'SIG', 'PH TA106 02'), [
            _mot(1, 2, 'A2'), _mot(31, 2, 'PH'), _mot(37, 2, 'A104'), _mot(47, 2, '02'),
            _mot(61, 2, 'SIG'), _mot(91, 2, 'PH'), _mot(97, 2, 'TA106'), _mot(107, 2, '02')])
        self.assertEqual(row['debuts'], {0: [0], 1: [0, 6, 16], 2: [0], 3: [0, 6, 16]})
        self.assertEqual((bilan['par_rang'], bilan['sans_jumeau']), (0, 0))

    def test_retrait_d_une_cellule_garde(self):
        row, _ = self._placer(_ligne('A2', 'RM 03B', '', ''),
                              [_mot(1, 2, 'A2'), _mot(31, 2, 'RM'), _mot(42, 2, '03B')])
        self.assertEqual(row['debuts'][1], [0, 11])

    def test_replis_o_0_et_i_1(self):
        row, bilan = self._placer(_ligne('A2', 'PH A104 02', '', ''), [
            _mot(1, 2, 'A2'), _mot(31, 2, 'PH'), _mot(37, 2, 'AIO4'), _mot(47, 2, 'O2')])
        self.assertEqual(row['debuts'][1], [0, 6, 16])
        self.assertEqual(bilan['par_rang'], 0)

    def test_jumeau_par_rang_dans_la_cellule(self):
        row, bilan = self._placer(_ligne('A2', 'PH A104 02', '', ''), [
            _mot(1, 2, 'A2'), _mot(31, 2, 'FH'), _mot(37, 2, 'A104'), _mot(47, 2, '02')])
        self.assertEqual(row['debuts'][1], [0, 6, 16])
        self.assertEqual(bilan['par_rang'], 1)

    def test_par_rang_garde_fou_un_espace(self):
        row, _ = self._placer(_ligne('A2', 'ABCD EF', 'SIGNAL LONG', 'PH TA106 02'), [
            _mot(1, 2, 'A2'), _mot(31, 2, 'XY'), _mot(33, 2, 'ZT'), _mot(61, 2, 'SIGNAL'),
            _mot(68, 2, 'LONG'), _mot(91, 2, 'PH'), _mot(97, 2, 'TA106'), _mot(107, 2, '02')])
        self.assertEqual(row['debuts'][1], [0, 5])

    def test_sans_jumeau_un_espace_apres_le_precedent(self):
        row, bilan = self._placer(_ligne('A2', 'PH A104 02', '', ''), [
            _mot(1, 2, 'A2'), _mot(31, 2, 'PH'), _mot(37, 2, 'A104')])
        self.assertEqual(row['debuts'][1], [0, 6, 11])
        self.assertEqual(bilan['sans_jumeau'], 1)

    def test_premier_mot_sans_jumeau_en_colonne_0(self):
        row, _ = self._placer(_ligne('A2', 'ZZ A104', '', ''),
                              [_mot(1, 2, 'A2'), _mot(37, 2, 'A104'), _mot(40, 2, 'QQ')])
        self.assertEqual(row['debuts'][1], [0, 6])

    def test_jumeau_dans_la_colonne_voisine_deplace(self):
        row, bilan = self._placer(_ligne('A2', 'PH A104', '02 SIG', ''), [
            _mot(1, 2, 'A2'), _mot(31, 2, 'PH'), _mot(37, 2, 'A104'), _mot(47, 2, '02'),
            _mot(61, 2, 'SIG')])
        self.assertEqual(row['cells'][1:3], ['PH A104 02', 'SIG'])
        self.assertEqual(row['debuts'][1], [0, 6, 16])
        self.assertEqual(len(bilan['deplaces']), 1)
        self.assertEqual(bilan['deplaces'][0]['mot'], '02')
        self.assertEqual((bilan['deplaces'][0]['de'], bilan['deplaces'][0]['vers']),
                         ('SIGNAL', 'TENANT'))

    def test_jumeau_de_la_cellule_avant_celui_de_la_voisine(self):
        # 223111PE011 p. 7, ligne 4 : « PH » de fin de SIGNAL lu « FH », « PH » d'ABOUTISSANT
        # bien lu ; le « PH » de SIGNAL prend la place de « FH » (rang), il ne bouge pas.
        row, bilan = self._placer(_ligne('A2', 'PH A104 02', 'ZONE D PH', 'PH TB203 02'), [
            _mot(1, 2, 'A2'), _mot(31, 2, 'PH'), _mot(37, 2, 'A104'), _mot(47, 2, '02'),
            _mot(61, 2, 'ZONE'), _mot(66, 2, 'D'), _mot(68, 2, 'FH'), _mot(91, 2, 'PH'),
            _mot(97, 2, 'T2203'), _mot(107, 2, '02')])
        self.assertEqual(row['cells'][2:], ['ZONE D PH', 'PH TB203 02'])
        self.assertEqual(row['debuts'][2], [0, 5, 7])
        self.assertEqual(bilan['deplaces'], [])

    def test_debordement_vers_la_gauche_pas_deplace(self):
        # 223111PE011 p. 4, ligne 1 : « PH » de TENANT imprimé à cheval sur le trait (fin à
        # moins d'un caractère du trait) ; il reste dans TENANT.
        row, bilan = self._placer(_ligne('N', 'PH AD/BNC 03', 'SIG', ''), [
            _mot(13, 2, 'N'), _mot(28, 2, 'PH'), _mot(35, 2, 'AVEC'), _mot(47, 2, '03'),
            _mot(61, 2, 'SIG')])
        self.assertEqual(row['cells'][:2], ['N', 'PH AD/BNC 03'])
        self.assertEqual(bilan['deplaces'], [])

    def test_mot_au_milieu_jamais_deplace(self):
        row, bilan = self._placer(_ligne('A2', 'PH', 'X 02 Y', ''), [
            _mot(1, 2, 'A2'), _mot(31, 2, 'PH'), _mot(47, 2, '02'), _mot(61, 2, 'X'),
            _mot(70, 2, 'Y')])
        self.assertEqual(row['cells'][1:3], ['PH', 'X 02 Y'])
        self.assertEqual(bilan['deplaces'], [])

    def test_jumeau_a_deux_colonnes_pas_deplace(self):
        row, bilan = self._placer(_ligne('A2', 'PH', '', 'TA106'), [
            _mot(1, 2, 'A2'), _mot(31, 2, 'PH'), _mot(37, 2, 'TA106')])
        self.assertEqual(row['cells'][1:], ['PH', '', 'TA106'])
        self.assertEqual(bilan['deplaces'], [])

    def test_contenu_inchange(self):
        claude = _ligne('A2', 'PH A104', '02 SIG', 'PH TA106 02')
        row, _ = self._placer(claude, [
            _mot(1, 2, 'A2'), _mot(31, 2, 'FH'), _mot(47, 2, '02'), _mot(61, 2, 'SIG')])
        self.assertEqual(_mots(row), _mots(claude))

    def test_section_une_seule_cellule(self):
        rows, _ = placer_page(
            [REPERE_CLAUDE, {'type': 'section', 'text': 'NOM DU CABLE : WPHR/A105'}],
            COLONNES, _lecture(*REPERE_TESS, _mot(1, 2, 'NOM'), _mot(5, 2, 'DU'),
                               _mot(8, 2, 'CABLE'), _mot(14, 2, ':'), _mot(19, 2, 'WPHR/A105')))
        self.assertEqual(rows[1]['debuts'], {0: [0, 4, 7, 13, 18]})

    def test_lignes_appariees_malgre_bruit_et_ligne_absente(self):
        claude = [REPERE_CLAUDE, _ligne('A2', 'QQ 1', '', ''), _ligne('A3', 'RR 2', '', '')]
        tess = [_mot(40, 2, '|'), _mot(1, 3, 'A3'), _mot(31, 3, 'RR'), _mot(36, 3, '2')]
        rows, bilan = placer_page(claude, COLONNES, _lecture(*REPERE_TESS, *tess))
        self.assertNotIn('debuts', rows[1])
        self.assertEqual(rows[2]['debuts'][1], [0, 5])
        self.assertEqual(bilan['lignes_sans_partenaire'], 1)

    def test_entree_non_modifiee(self):
        import copy
        claude = [REPERE_CLAUDE, _ligne('A2', 'PH A104', '02 SIG', '')]
        avant = copy.deepcopy(claude)
        placer_page(claude, COLONNES, _lecture(*REPERE_TESS, _mot(47, 2, '02')))
        self.assertEqual(claude, avant)

    def test_lecture_vers_dict_et_retour(self):
        lecture = _lecture(*REPERE_TESS)
        self.assertEqual(LecturePage.depuis_dict(lecture.vers_dict()), lecture)


class TestTexteAuxPositions(unittest.TestCase):

    def test_complete_par_des_espaces(self):
        self.assertEqual(texte_aux_positions('PH A104 02', [0, 6, 16]), 'PH    A104      02')

    def test_retrait(self):
        self.assertEqual(texte_aux_positions('RM 03B', [1, 12]), ' RM         03B')

    def test_nombre_de_mots_different_inchange(self):
        self.assertEqual(texte_aux_positions('PH A104', [0, 6, 16]), 'PH A104')


class TestExcel(unittest.TestCase):

    def _excel(self, actives=True):
        tpl = TemplateManager().get('REPARTITEUR 2')
        page = {'success': True, 'headers': COLONNES, 'image_path': 'page_052.png',
                'detection_method': 'claude-vision', 'metadata': {'PAGE': '52'},
                'rows': [dict(_ligne('A2', 'PH A104 02', 'SIG', ''),
                              debuts={1: [0, 6, 16]}),
                         _ligne('A3', 'PH A104 03', 'SIG', '')]}
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(Config, 'POSITIONS_ORIGINALES', actives):
            sortie = Path(tmp) / 's.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                generer_excel([page], PdfTableExtractor(tpl), sortie, on_log=lambda m: None)
            ws = openpyxl.load_workbook(sortie).worksheets[0]
            return [(ws.cell(row=r, column=2).value, ws.cell(row=r, column=2).font)
                    for r in (2, 3)]

    def test_sous_champs_a_leur_colonne_et_courier_new_11(self):
        (v2, f2), (v3, f3) = self._excel()
        self.assertEqual(v2, 'PH    A104      02')
        self.assertEqual(v3, 'PH A104 03')
        for police in (f2, f3):
            self.assertEqual((police.name, police.sz), (Config.POSITIONS_POLICE,
                                                        Config.POSITIONS_TAILLE_POLICE))

    def test_desactive_comportement_actuel(self):
        (v2, f2), _ = self._excel(actives=False)
        self.assertEqual(v2, 'PH A104 02')
        self.assertNotEqual(f2.name, Config.POSITIONS_POLICE)


if __name__ == '__main__':
    unittest.main()
