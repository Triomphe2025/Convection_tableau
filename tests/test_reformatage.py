"""
Reformatage (« Formater Excel », generer_classeur.reformatter_excel) des Excel produits par la
version actuelle : un extrait par modèle de sortie, converti en rejeu des réponses Claude
enregistrées (aucun appel API), puis reformaté.

Contrôles : même nombre de tableaux ; chaque bloc a la taille du plus grand tableau ; les 2
lignes du pied en bas de chaque bloc ; valeurs, couleurs et commentaires conservés ; chaque
lien d'A VERIFIER pointe sur la même valeur qu'avant ; LEGENDE recopiée telle quelle.
Régression évitée : le logo imprimé « M A T R A » et le pied « CABLE : … TYPE : … » n'étaient
plus reconnus (blocs de 118 lignes, pied coupé en deux), A VERIFIER était prise pour des
tableaux (liens perdus), commentaires et liens n'étaient pas copiés.

Lancement :
    env\\Scripts\\python.exe -m pytest tests/test_reformatage.py -v
"""
import contextlib
import io
import json
import shutil
import tempfile
import unittest
from collections import Counter
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl

from config import Config
from converter import Converter
from generer_classeur import FEUILLE_A_VERIFIER, FEUILLE_LEGENDE, reformatter_excel
from pied_page import ligne_de_pied
from template import TemplateManager

FIX = Path(__file__).parent / 'fixtures'
EXTRAITS = [('223111PE011_extrait_10pages', 'REPARTITEUR 2'),
            ('6A23111PE133_extrait_8pages', 'REPARTITEUR'),
            ('223111PE012_extrait_10pages', 'Bornier standard')]
MOTS_ENTETE = {c.upper() for n in TemplateManager().names()
               for c in TemplateManager().get(n).columns}


def _faux_client(journal: Path):
    reponses = [json.loads(lg)['raw'] for lg in journal.read_text('utf-8').splitlines()
                if lg.strip()]

    class Faux:
        def __init__(self, api_key=None, **kw):
            self.messages = SimpleNamespace(create=self._create)

        def _create(self, **kw):
            return SimpleNamespace(
                id='msg_rejeu', content=[SimpleNamespace(type='text', text=reponses.pop(0))],
                stop_reason='end_turn', model=kw['model'],
                usage=SimpleNamespace(input_tokens=0, output_tokens=0,
                                      cache_read_input_tokens=0,
                                      cache_creation_input_tokens=0))
    return Faux


def _convertir(stem, gabarit, dossier: Path) -> Path:
    journal = FIX / 'rejeu_claude' / f'{stem}_claude.jsonl'
    with patch.object(Config, 'OCR_MODE', 'claude'), \
            patch.object(Config, 'CLAUDE_API_KEY', 'cle-factice'), \
            patch('anthropic.Anthropic', _faux_client(journal)), \
            contextlib.redirect_stdout(io.StringIO()):
        Converter(word_file=FIX / f'{stem}.pdf', output_dir=dossier,
                  template=TemplateManager().get(gabarit), on_log=lambda m: None).run()
    return dossier / f'{stem}.xlsx'


def _entetes(ws):
    return [r[0].row for r in ws.iter_rows()
            if sum(1 for c in r if str(c.value or '').strip().upper() in MOTS_ENTETE) >= 2]


def _vide(ligne):
    return not any(str(c.value or '').strip() for c in ligne)


def _tailles_tableaux(ws):
    """Lignes utiles de chaque tableau : de l'en-tête à la dernière ligne du pied."""
    lignes = list(ws.iter_rows())
    entetes = [r - 1 for r in _entetes(ws)] + [len(lignes)]
    tailles = []
    for debut, fin in zip(entetes, entetes[1:]):
        bloc = lignes[debut:fin]
        pieds = [i for i, lg in enumerate(bloc) if ligne_de_pied(c.value for c in lg)]
        donnees = [i for i, lg in enumerate(bloc) if not _vide(lg) and i not in pieds]
        tailles.append(max(donnees) + 1 + len(pieds))
    return tailles


def _cellules(ws):
    return [c for ligne in ws.iter_rows() for c in ligne if str(c.value or '').strip()]


def _cible(wb, lien):
    feuille, coord = lien.split('!')
    return wb[feuille.strip("'")][coord].value


class TestReformatage(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls._tmp = Path(tempfile.mkdtemp(prefix='reformatage_'))
        cls.cas = {}
        for stem, gabarit in EXTRAITS:
            dossier = cls._tmp / stem
            source = _convertir(stem, gabarit, dossier)
            sortie = dossier / 'reformate.xlsx'
            journal = []
            n = reformatter_excel(source, sortie, on_log=journal.append)
            cls.cas[stem] = (openpyxl.load_workbook(source), openpyxl.load_workbook(sortie),
                             n, journal)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls._tmp, ignore_errors=True)

    def _chaque(self):
        for stem, cas in self.cas.items():
            with self.subTest(extrait=stem):
                yield cas

    def test_meme_nombre_de_tableaux(self):
        for avant, apres, n, _ in self._chaque():
            self.assertEqual(len(_entetes(avant['Borniers'])), n)
            self.assertEqual(len(_entetes(apres['Borniers'])), n)

    def test_blocs_a_la_taille_du_plus_grand_tableau(self):
        for avant, apres, _, _ in self._chaque():
            entetes = _entetes(apres['Borniers'])
            blocs = {b - a for a, b in zip(entetes, entetes[1:])}
            attendu = max(Config.PAGE_SIZE, max(_tailles_tableaux(avant['Borniers'])))
            self.assertEqual(blocs, {attendu})

    def test_deux_lignes_de_pied_en_bas_de_chaque_bloc(self):
        for _, apres, _, _ in self._chaque():
            ws = apres['Borniers']
            entetes = _entetes(ws)
            taille = entetes[1] - entetes[0]
            for debut in entetes:
                fin = debut + taille - 1
                pied = [ligne_de_pied(c.value for c in ws[r]) for r in (fin - 2, fin - 1, fin)]
                self.assertEqual(pied, [False, True, True], f"bloc ligne {debut}")

    def test_valeurs_couleurs_commentaires_conserves(self):
        for avant, apres, _, _ in self._chaque():
            a, b = _cellules(avant['Borniers']), _cellules(apres['Borniers'])
            self.assertEqual(Counter(c.value for c in a), Counter(c.value for c in b))
            couleur = (lambda c: (c.value, str(c.fill.fgColor.rgb))  # noqa: E731
                       if c.fill.fill_type == 'solid' else None)
            self.assertEqual(Counter(map(couleur, a)), Counter(map(couleur, b)))
            commentaire = (lambda c: (c.value, c.comment.text)  # noqa: E731
                           if c.comment else None)
            self.assertEqual(Counter(map(commentaire, a)), Counter(map(commentaire, b)))

    def test_liens_d_a_verifier_sur_la_meme_valeur(self):
        verifies = 0
        for avant, apres, _, _ in self._chaque():
            if FEUILLE_A_VERIFIER not in avant.sheetnames:
                continue
            lignes_a = list(avant[FEUILLE_A_VERIFIER].iter_rows())
            lignes_b = list(apres[FEUILLE_A_VERIFIER].iter_rows())
            self.assertEqual([[c.value for c in r[:6]] for r in lignes_a],
                             [[c.value for c in r[:6]] for r in lignes_b])
            for ra, rb in zip(lignes_a, lignes_b):
                if ra[-1].hyperlink is None:
                    continue
                self.assertEqual(_cible(apres, rb[-1].hyperlink.location),
                                 _cible(avant, ra[-1].hyperlink.location))
                self.assertEqual(rb[-1].value, '→ ' + rb[-1].hyperlink.location.split('!')[1])
                verifies += 1
        self.assertGreater(verifies, 0)

    def test_legende_recopiee_telle_quelle(self):
        for avant, apres, _, _ in self._chaque():
            if FEUILLE_LEGENDE not in avant.sheetnames:
                continue
            self.assertEqual(
                [(c.value, str(c.fill.fgColor.rgb)) for c in _cellules(avant[FEUILLE_LEGENDE])],
                [(c.value, str(c.fill.fgColor.rgb)) for c in _cellules(apres[FEUILLE_LEGENDE])])

    def test_aucun_avertissement_sur_les_extraits(self):
        for _, _, _, journal in self._chaque():
            self.assertEqual([m for m in journal if '⚠' in m], [])

    def test_ordre_des_feuilles_conserve(self):
        for avant, apres, _, _ in self._chaque():
            self.assertEqual(avant.sheetnames, apres.sheetnames)


class TestLignesDePied(unittest.TestCase):

    def test_logo_espace_et_cable_type(self):
        self.assertTrue(ligne_de_pied(['M A T R A', 'CABLE : WPHA104/TA106  TYPE : 30P887']))
        self.assertTrue(ligne_de_pied(['', 'CABLE : X  TYPE : Y']))
        self.assertTrue(ligne_de_pied(['N°PLAN : 223111PE011  | INDICE : R  | PAGE : 52']))

    def test_section_et_donnees_pas_un_pied(self):
        self.assertFalse(ligne_de_pied(['NOM DU CABLE : WPHR/SON01']))
        self.assertFalse(ligne_de_pied(['01', 'TS MTI DEFAUT', 'RESERVE']))
        self.assertFalse(ligne_de_pied(['CABLE : X']))


if __name__ == '__main__':
    unittest.main()


class TestAvertissements(unittest.TestCase):

    def test_pied_court_signale(self):
        from generer_classeur import avertissements_reformatage
        alertes = avertissements_reformatage([(40, 2), (40, 1), (30, 2), (40, 1)], 59, 59)
        self.assertEqual(alertes, [
            "pied reconnu sur 1 ligne(s) au lieu de 2 : tableaux 2 et 4 ; logo inconnu ? voir "
            "la liste des logos dans config.py (PIED_LOGOS)"])

    def test_bloc_trop_grand_signale(self):
        from generer_classeur import avertissements_reformatage
        alertes = avertissements_reformatage([(58, 2), (116, 2)], 118, 59)
        self.assertEqual(len(alertes), 1)
        self.assertIn('bloc uniforme de 118 lignes', alertes[0])
        self.assertIn('tableau 2 ;', alertes[0])

    def test_rien_a_signaler(self):
        from generer_classeur import avertissements_reformatage
        self.assertEqual(avertissements_reformatage([(40, 2), (55, 2)], 59, 59), [])

    def test_logo_inconnu_au_journal_de_la_page_formater(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Borniers'
        for logo in ('SIEMENS', 'A C M E'):
            ws.append(['FIL', 'TENANT', 'SIGNAL', 'ABOUTISSANT'])
            for k in range(3):
                ws.append([str(k), f'PH A104 0{k}', 'SIG', 'RA B1 17'])
            ws.append([logo, 'REF CE 8707905'])
            ws.append(['N°PLAN : 223111PE011', '', 'INDICE : R', 'PAGE : 1'])
        source = Path(tempfile.mkdtemp()) / 's.xlsx'
        wb.save(source)
        journal = []
        reformatter_excel(source, source.with_name('r.xlsx'), on_log=journal.append)
        self.assertIn("  ⚠ pied reconnu sur 1 ligne(s) au lieu de 2 : tableau 2 ; logo inconnu ? "
                      "voir la liste des logos dans config.py (PIED_LOGOS)", journal)
