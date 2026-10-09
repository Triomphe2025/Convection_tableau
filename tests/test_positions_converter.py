"""
Commit A dans le convertisseur : positions d'origine posées sur les pages lues par
Claude (faux client, réponse enregistrée, Tesseract en local), lecture Tesseract gardée
dans le cache de l'appli (jamais à côté des fichiers de l'utilisateur), réutilisée sans
2e OCR ; sans PDF (rejeu d'un journal), une ligne au journal.

Lancement :
    env\\Scripts\\python.exe -m pytest tests/test_positions_converter.py -v
"""
import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import fitz
import openpyxl

from config import Config
from converter import Converter
import positions_scan
from positions_scan import LecturePage
from template import TemplateManager

FIX = Path(__file__).parent / 'fixtures'
REPONSE_P15 = next(json.loads(lg)['raw'] for lg in
                   (FIX / 'positions_reponses_claude.jsonl').read_text('utf-8').splitlines()
                   if json.loads(lg)['document'].startswith('6A23111PE133'))


class _FauxClient:
    def __init__(self, api_key=None):
        self.messages = SimpleNamespace(create=self._create)

    def _create(self, **kwargs):
        return SimpleNamespace(
            id='msg_faux', content=[SimpleNamespace(type='text', text=REPONSE_P15)],
            stop_reason='end_turn', model=kwargs['model'],
            usage=SimpleNamespace(input_tokens=1, output_tokens=1, cache_read_input_tokens=0,
                                  cache_creation_input_tokens=0))


@contextlib.contextmanager
def _mode_claude():
    noms = ('OCR_MODE', 'CLAUDE_API_KEY')
    avant = {n: getattr(Config, n) for n in noms}
    Config.OCR_MODE, Config.CLAUDE_API_KEY = 'claude', 'cle-factice'
    try:
        yield
    finally:
        for n, v in avant.items():
            setattr(Config, n, v)


@unittest.skipUnless(Path(Config.TESSERACT_PATH).exists(), "Tesseract absent")
class TestConversionPdf(unittest.TestCase):

    @classmethod
    def _convertir(cls, sortie):
        journal = []
        with _mode_claude(), patch('anthropic.Anthropic', _FauxClient), \
                patch.object(Config, 'POSITIONS_CACHE_DOSSIER', cls.cache), \
                contextlib.redirect_stdout(io.StringIO()):
            resultat = Converter(word_file=cls.pdf, output_dir=cls.dossier / sortie,
                                 template=TemplateManager().get('REPARTITEUR'),
                                 on_log=journal.append).run()
        return resultat, journal

    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        cls.dossier = Path(cls._tmp.name)
        cls.cache = cls.dossier / 'cache_appli'
        cls.pdf = cls.dossier / 'pe133_p15.pdf'
        with fitz.open(str(FIX / '6A23111PE133_extrait_8pages.pdf')) as src, fitz.open() as un:
            un.insert_pdf(src, from_page=5, to_page=5)
            un.save(str(cls.pdf))
        cls.resultat, cls.journal = cls._convertir('sortie')
        ws = openpyxl.load_workbook(cls.resultat['excel']).worksheets[0]
        cls.tenants = [ws.cell(row=r, column=1).value for r in range(2, 58)]

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_bilan_au_journal(self):
        lignes = [m for m in self.journal if "positions d'origine p. 1 :" in m]
        self.assertEqual(len(lignes), 1)
        self.assertIn('0 déplacé(s)', lignes[0])

    def test_positions_dans_l_excel(self):
        self.assertIn('P111TC    26', self.tenants)
        self.assertIn('RM         03B', self.tenants)

    def test_lecture_dans_le_cache_de_l_appli_pas_a_cote_des_fichiers(self):
        fichiers = list(self.cache.glob('*.json'))
        self.assertEqual(len(fichiers), 1)
        donnees = json.loads(fichiers[0].read_text('utf-8'))
        self.assertGreater(len(LecturePage.depuis_dict(donnees).mots), 100)
        self.assertEqual(list((self.dossier / 'sortie').rglob('*.json')), [])

    def test_deuxieme_conversion_sans_deuxieme_ocr(self):
        with patch.object(positions_scan, 'lire_page',
                          side_effect=AssertionError('2e OCR')) as lire:
            _, journal = self._convertir('sortie_2')
        lire.assert_not_called()
        self.assertTrue([m for m in journal if "positions d'origine p. 1 :" in m])


@unittest.skipUnless(Path(Config.TESSERACT_PATH).exists(), "Tesseract absent")
class TestTesseractPendantClaude(unittest.TestCase):
    """Tesseract lit les pages scannées dans un fil à part pendant que Claude les lit."""

    def test_lecture_dans_un_fil_a_part_meme_excel(self):
        import threading
        import time
        evenements = []
        vraie_lecture = positions_scan.lire_page

        def lecture_suivie(page):
            evenements.append(('tesseract_debut', threading.current_thread().name))
            resultat = vraie_lecture(page)
            evenements.append(('tesseract_fin', None))
            return resultat

        class ClaudeLent(_FauxClient):
            def _create(self, **kwargs):
                evenements.append(('claude_debut', None))
                time.sleep(3)          # plus long que la lecture Tesseract de la page
                evenements.append(('claude_fin', None))
                return super()._create(**kwargs)

        with tempfile.TemporaryDirectory() as tmp:
            dossier = Path(tmp)
            pdf = dossier / 'p15.pdf'
            with fitz.open(str(FIX / '6A23111PE133_extrait_8pages.pdf')) as src, \
                    fitz.open() as un:
                un.insert_pdf(src, from_page=5, to_page=5)
                un.save(str(pdf))
            with _mode_claude(), patch('anthropic.Anthropic', ClaudeLent), \
                    patch.object(positions_scan, 'lire_page', lecture_suivie), \
                    patch.object(Config, 'POSITIONS_CACHE_DOSSIER', dossier / 'cache'), \
                    contextlib.redirect_stdout(io.StringIO()):
                resultat = Converter(word_file=pdf, output_dir=dossier / 'sortie',
                                     template=TemplateManager().get('REPARTITEUR'),
                                     on_log=lambda m: None).run()
            ws = openpyxl.load_workbook(resultat['excel']).worksheets[0]
            tenants = [ws.cell(row=r, column=1).value for r in range(2, 58)]
        noms = [e[0] for e in evenements]
        self.assertTrue(evenements[noms.index('tesseract_debut')][1].startswith('tesseract'))
        # Tesseract a commencé avant que Claude ait rendu la page : les deux lectures se
        # recouvrent (la fin de chacune dépend de la charge de la machine, pas du code).
        self.assertLess(noms.index('tesseract_debut'), noms.index('claude_fin'))
        self.assertIn('P111TC    26', tenants)


@unittest.skipUnless(Path(Config.TESSERACT_PATH).exists(), "Tesseract absent")
class TestIndicateurDePage(unittest.TestCase):
    """Taux de divergence Claude / Tesseract au journal ; page dégradée = ligne ℹ et A VERIFIER."""

    def _convertir(self, document, page, gabarit):
        entrees = [json.loads(lg) for lg in
                   (FIX / 'positions_reponses_claude.jsonl').read_text('utf-8').splitlines()]
        reponse = next(e['raw'] for e in entrees
                       if e['document'] == document and e['page'] == page)

        class Client(_FauxClient):
            def _create(self, **kwargs):
                retour = super()._create(**kwargs)
                retour.content[0].text = reponse
                return retour

        journal = []
        with tempfile.TemporaryDirectory() as tmp:
            dossier = Path(tmp)
            pdf = dossier / 'une_page.pdf'
            with fitz.open(str(FIX / document)) as src, fitz.open() as un:
                un.insert_pdf(src, from_page=page - 1, to_page=page - 1)
                un.save(str(pdf))
            with _mode_claude(), patch('anthropic.Anthropic', Client), \
                    patch.object(Config, 'POSITIONS_CACHE_DOSSIER', dossier / 'cache'), \
                    contextlib.redirect_stdout(io.StringIO()):
                resultat = Converter(word_file=pdf, output_dir=dossier / 'sortie',
                                     template=TemplateManager().get(gabarit),
                                     on_log=journal.append).run()
            wb = openpyxl.load_workbook(resultat['excel'])
            feuille = ([[c.value for c in r] for r in wb['A VERIFIER'].iter_rows()]
                       if 'A VERIFIER' in wb.sheetnames else None)
        return journal, feuille

    def test_page_propre_taux_au_journal_sans_a_verifier(self):
        journal, feuille = self._convertir('6A23111PE133_extrait_8pages.pdf', 6, 'REPARTITEUR')
        bilan = next(m for m in journal if "positions d'origine p. 1 :" in m)
        self.assertRegex(bilan, r'divergence Claude / Tesseract \d+ %')
        self.assertFalse([m for m in journal if 'ℹ' in m])
        # Page propre, réponse non modifiée : aucune alerte de conservation.
        self.assertFalse([m for m in journal if 'peut-être' in m])
        self.assertIsNone(feuille)

    def test_page_degradee_ligne_d_information_et_a_verifier(self):
        journal, feuille = self._convertir('223111PE011_extrait_10pages.pdf', 5,
                                           'REPARTITEUR 2')
        infos = [m.strip() for m in journal if 'ℹ' in m]
        self.assertEqual(len(infos), 1)
        self.assertRegex(infos[0], r'^ℹ p\. 1 : scan dégradé \(divergence \d+ %\) : à relire '
                                   r'en priorité ; contrôle de conservation impossible$')
        self.assertFalse([m for m in journal if '⚠' in m and 'scan' in m])
        # Une seule page scannée, dégradée : plus de la moitié → synthèse puis la page.
        self.assertEqual([r[3] for r in feuille[1:]], ['synthèse', 'page dégradée'])


class TestRejeuSansPdf(unittest.TestCase):

    def test_une_ligne_au_journal(self):
        journal = []
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / 'x_claude.jsonl'
            source.write_text(json.dumps({'image': 'page_006.png', 'success': True,
                                          'raw': REPONSE_P15}) + '\n', encoding='utf-8')
            with contextlib.redirect_stdout(io.StringIO()):
                Converter(word_file=source, output_dir=Path(tmp) / 'sortie',
                          template=TemplateManager().get('REPARTITEUR'),
                          on_log=journal.append).run()
        self.assertIn("  positions d'origine non recalculées : rejeu sans le PDF", journal)


if __name__ == '__main__':
    unittest.main()
