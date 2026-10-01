"""
Champs constants d'un document (P.E.T. par défaut) : une case vide est reprise
des autres pages du MÊME document si elles concordent toutes, colorée et commentée.
Une valeur lue n'est jamais remplacée ; INDICE et PAGE ne sont jamais repris.

Lancement :
    env\\Scripts\\python.exe -m pytest tests/test_champs_constants.py -v
"""
import contextlib
import copy
import io
import tempfile
import unittest
from pathlib import Path

import fitz
import openpyxl

import pdf_extractor as pe
from config import Config
from generer_classeur import deduire_champs_constants, generer_excel, marquer_deduits
from pied_page import inserer_valeur
from template import TableTemplate, TemplateManager

PE137 = Path(__file__).parent / 'fixtures' / '223400PE137.pdf'
DEDUIT = "déduit des autres pages du document"


def _page(page, **meta):
    return {'success': True, 'headers': ['FIL'], 'image_path': f'page_{page}.png',
            'metadata': dict(PAGE=page, **meta),
            'rows': [{'type': 'data', 'cells': ['X'], 'confidence': [100]}]}


class TestPE137UnPetVide(unittest.TestCase):
    """P.E.T. = GRAND-BUT sur 47 pages, vidé sur 1 : cette page reçoit GRAND-BUT, colorée."""

    @classmethod
    def setUpClass(cls):
        tpl = TemplateManager().get('REPARTITEUR')
        lecteur = pe.PdfTableExtractor(tpl)
        with fitz.open(str(PE137)) as doc:
            resultats = [lecteur.extract_page_grille(doc[i], i) for i in range(len(doc))]
        tableaux = [r for r in resultats if r['success']]
        cls.cible = tableaux[10]
        meta = cls.cible['metadata']
        cls.page = meta['PAGE']
        del meta['PET']
        meta['PIED_BRUT'] = [ligne.replace('GRAND-BUT', '         ')
                             for ligne in meta['PIED_BRUT']]
        cls.messages = []
        with tempfile.TemporaryDirectory() as tmp:
            sortie = Path(tmp) / 's.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                generer_excel(resultats, lecteur, sortie, on_log=cls.messages.append)
            ws = openpyxl.load_workbook(sortie).worksheets[0]
            cls.pieds = [ws.cell(row=c.row - 1, column=2) for c in ws['B']
                         if isinstance(c.value, str) and 'INDICE' in c.value]

    def test_48_pieds(self):
        self.assertEqual(len(self.pieds), 48)

    def test_page_videe_recoit_grand_but_coloree_et_commentee(self):
        cellule = self.pieds[10]
        self.assertIn('GRAND-BUT', cellule.value)
        self.assertEqual(cellule.fill.fgColor.rgb, '00' + Config.COULEUR_DEDUIT)
        self.assertIn(DEDUIT, cellule.comment.text)

    def test_autres_pages_ni_colorees_ni_commentees(self):
        autres = [c for i, c in enumerate(self.pieds) if i != 10]
        self.assertTrue(all(c.comment is None for c in autres))
        self.assertTrue(all(c.fill.fill_type is None for c in autres))

    def test_pas_d_alerte_pet_absent(self):
        self.assertFalse([m for m in self.messages if 'PET' in m])

    def test_resultat_d_origine_non_modifie(self):
        self.assertNotIn('PET', self.cible['metadata'])


class TestDeduireChampsConstants(unittest.TestCase):

    def test_valeurs_differentes_case_vide_et_alerte(self):
        pages = [_page('1', PET='GRAND-BUT'), _page('2', PET='EPEULE'), _page('3')]
        resultats, alertes = deduire_champs_constants(pages, ['PET'], {})
        self.assertNotIn('PET', resultats[2]['metadata'])
        self.assertEqual(alertes, ["PET vide, non complété : page(s) 3 — les autres pages "
                                   "ne concordent pas (GRAND-BUT, EPEULE)"])

    def test_valeur_lue_jamais_remplacee(self):
        pages = [_page('1', PET='GRAND-BUT'), _page('2', PET='GRAND-BUT'),
                 _page('3', PET='GRAND BUT')]
        resultats, alertes = deduire_champs_constants(pages, ['PET'], {})
        self.assertEqual(resultats[2]['metadata']['PET'], 'GRAND BUT')
        self.assertEqual(alertes, [])

    def test_indice_vide_jamais_complete(self):
        pages = [_page('1', INDICE='R'), _page('2', INDICE='R'), _page('3')]
        resultats, alertes = deduire_champs_constants(pages, ['INDICE', 'PAGE'], {})
        self.assertNotIn('INDICE', resultats[2]['metadata'])
        self.assertEqual(alertes, [])

    def test_champ_absent_partout_rien_a_reprendre(self):
        resultats, alertes = deduire_champs_constants([_page('1'), _page('2')], ['PET'], {})
        self.assertNotIn('PET', resultats[0]['metadata'])
        self.assertEqual(alertes, [])

    def test_trace_de_la_deduction_et_pied_brut(self):
        pages = [_page('1', PET='GRAND-BUT'),
                 _page('2', PIED_BRUT=['P.E.T. :      JARRETIERAGE', 'N° PLAN : P'])]
        resultats, _ = deduire_champs_constants(pages, ['PET'], {'PET': 'P.E.T.'})
        meta = resultats[1]['metadata']
        self.assertEqual((meta['PET'], meta['DEDUITS']), ('GRAND-BUT', ['PET']))
        self.assertEqual(meta['PIED_BRUT'][0], 'P.E.T. : GRAND-BUT      JARRETIERAGE')

    def test_entree_non_modifiee(self):
        pages = [_page('1', PET='GRAND-BUT'), _page('2')]
        avant = copy.deepcopy(pages)
        deduire_champs_constants(pages, ['PET'], {})
        self.assertEqual(pages, avant)


class TestInsererValeur(unittest.TestCase):

    def test_apres_le_libelle_vide(self):
        self.assertEqual(inserer_valeur(['CABLE : X', 'P.E.T. :   JARRETIERAGE'], 'PET',
                                        'P.E.T.', 'GRAND-BUT'),
                         ['CABLE : X', 'P.E.T. : GRAND-BUT   JARRETIERAGE'])

    def test_libelle_absent_ajoute_en_fin_de_premiere_ligne(self):
        self.assertEqual(inserer_valeur(['CABLE : X', 'N° PLAN : P'], 'PET', 'P.E.T.', 'G'),
                         ['CABLE : X     P.E.T. : G', 'N° PLAN : P'])

    def test_sans_ligne(self):
        self.assertEqual(inserer_valeur([], 'PET', 'P.E.T.', 'G'), ['P.E.T. : G'])


class TestMarquerDeduits(unittest.TestCase):

    def test_colore_la_ligne_qui_porte_la_valeur(self):
        ws = openpyxl.Workbook().active
        ws.cell(row=8, column=2, value='P.E.T. : GRAND-BUT   JARRETIERAGE')
        ws.cell(row=9, column=2, value='N° PLAN : P  INDICE : R')
        n = marquer_deduits(ws, 10, {'PET': 'GRAND-BUT', 'DEDUITS': ['PET']}, 'DDEBF7')
        self.assertEqual(n, 1)
        self.assertIn(DEDUIT, ws.cell(row=8, column=2).comment.text)
        self.assertIsNone(ws.cell(row=9, column=2).comment)

    def test_rien_de_deduit(self):
        ws = openpyxl.Workbook().active
        self.assertEqual(marquer_deduits(ws, 10, {'PET': 'X'}, 'DDEBF7'), 0)


class TestReglageParModele(unittest.TestCase):

    def test_pet_par_defaut(self):
        self.assertEqual(TableTemplate(name='t', columns=['A']).champs_constants, ['PET'])

    def test_ancien_templates_json_sans_le_reglage(self):
        tpl = TableTemplate.from_dict({'name': 't', 'columns': ['A']})
        self.assertEqual(tpl.champs_constants, ['PET'])

    def test_saisie_de_l_interface(self):
        self.assertEqual(TableTemplate.lire_champs_constants('P.E.T., N° PLAN, pet, , '),
                         ['PET', 'NO_PLAN'])
        self.assertEqual(TableTemplate.lire_champs_constants(''), [])


if __name__ == '__main__':
    unittest.main()
