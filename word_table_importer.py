"""
Importateur de tableaux depuis un document Word (.docx).

Lit chaque tableau d'un fichier Word déjà structuré (Document 2) et le
convertit en dict résultat compatible avec BornierTableExtractor._fill_worksheet().
Le numéro de PAGE extrait du pied de tableau est utilisé pour positionner
chaque tableau dans le classeur Excel principal (tri par page croissant).

Usage :
    from word_table_importer import WordTableImporter
    importer = WordTableImporter(extractor)
    results = importer.extract_tables(Path("tableaux_propres.docx"))
"""

import logging
from pathlib import Path
from typing import Dict, List, Optional

from docx import Document

logger = logging.getLogger(__name__)

# Mots-clés qui identifient une ligne de pied de page dans un tableau Word
_FOOTER_KEYWORDS = {
    'NO PLAN', 'INDICE', 'PAGE', 'P.E.T', 'BORNIER', 'M T I', 'M  T  I'
}


def _unique_cells(row):
    """Retourne les cellules uniques d'une ligne (ignore les doublons de fusion)."""
    seen: set = set()
    unique = []
    for cell in row.cells:
        cell_id = id(cell._tc)
        if cell_id not in seen:
            seen.add(cell_id)
            unique.append(cell)
    return unique


def _row_text(row) -> str:
    return ' '.join(c.text.strip() for c in _unique_cells(row))


def _is_footer_row(text: str) -> bool:
    upper = text.upper()
    return any(kw in upper for kw in _FOOTER_KEYWORDS)


class WordTableImporter:
    """
    Extrait les tableaux structurés d'un document Word (.docx).

    Chaque tableau est analysé pour en extraire :
      - la ligne d'en-tête (noms de colonnes)
      - les lignes de données
      - les métadonnées de pied (BORNIER, PAGE, P.E.T., INDICE, NO_PLAN)

    Le résultat est un dict identique à celui produit par
    BornierTableExtractor.extract() et compatible avec _fill_worksheet().
    """

    def __init__(self, extractor):
        """
        Args:
            extractor: Instance de BornierTableExtractor.
                       Réutilisé pour _extract_meta() et la cohérence du template.
        """
        self._extractor = extractor

    # ── API publique ──────────────────────────────────────────────────────

    def extract_tables(self, docx_path: Path) -> List[Dict]:
        """
        Extrait tous les tableaux d'un document Word.

        Args:
            docx_path: Chemin vers le fichier .docx

        Returns:
            Liste de dicts résultat triables par PAGE (clé 'metadata').
            Les tableaux sans lignes de données sont ignorés.
        """
        doc = Document(str(docx_path))
        total_tables = len(doc.tables)
        logger.info(
            f"WordTableImporter : {total_tables} tableau(x) brut(s)"
            f" dans {docx_path.name}"
        )

        results = []
        for idx, table in enumerate(doc.tables):
            n_rows = len(table.rows)
            n_cols_raw = len(table.columns)
            logger.debug(
                f"  Tableau {idx} : {n_rows} lignes × {n_cols_raw} colonnes"
            )
            result = self._parse_table(table, table_index=idx)
            if result is not None:
                meta = result.get('metadata', {})
                n_data = sum(
                    1 for r in result.get('rows', []) if r.get('type') == 'data'
                )
                logger.info(
                    f"  ✓ Tableau {idx} : {n_data} lignes de données"
                    f"  PAGE={meta.get('PAGE', '?')}"
                    f"  BORNIER={meta.get('BORNIER', '?')}"
                )
                results.append(result)
            else:
                logger.info(
                    f"  ✗ Tableau {idx} ignoré (vide, trop court ou sans données)"
                )

        logger.info(
            f"WordTableImporter : {len(results)}/{total_tables} tableau(x) valide(s)"
        )
        return results

    # ── Parsing interne ───────────────────────────────────────────────────

    def _parse_table(self, table, table_index: int) -> Optional[Dict]:
        """
        Convertit un objet Table python-docx en dict résultat.

        Returns None si le tableau est vide, trop court ou sans données.
        """
        rows = table.rows
        if len(rows) < 2:
            return None

        # ── Détection du début du pied (cherche depuis la fin, max 4 lignes) ──
        # On ne s'arrête pas au premier résultat : on continue tant que les
        # lignes sont des lignes de pied, pour inclure toutes les lignes du pied.
        footer_start = len(rows)
        for j in range(len(rows) - 1, max(len(rows) - 5, 0), -1):
            if _is_footer_row(_row_text(rows[j])):
                footer_start = j
            else:
                break

        # ── En-tête = première ligne ──────────────────────────────────────
        header_cells = [c.text.strip() for c in _unique_cells(rows[0])]
        if not any(header_cells):
            return None
        n_cols = len(header_cells)

        # ── Lignes de données ─────────────────────────────────────────────
        result_rows = []
        for row in rows[1:footer_start]:
            cells = [c.text.strip() for c in _unique_cells(row)]
            # Normaliser au nombre de colonnes de l'en-tête
            cells = (cells + [''] * n_cols)[:n_cols]
            if any(cells):
                result_rows.append({'type': 'data', 'cells': cells})

        if not result_rows:
            return None

        # ── Métadonnées depuis le pied ────────────────────────────────────
        footer_lines = []
        for row in rows[footer_start:]:
            text = _row_text(row)
            words = text.split()
            if words:
                footer_lines.append(
                    [{'text': w, 'cx': i * 20, 'cy': 0}
                     for i, w in enumerate(words)]
                )

        meta = self._extractor._extract_meta(footer_lines)

        return {
            'success': True,
            'headers': header_cells,
            'rows': result_rows,
            'metadata': meta,
            'image_path': f'word_table_{table_index}',
            'blur_pct': 0.0,  # tableaux Word : pas de flou d'image
        }
